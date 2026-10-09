import sys
import csv
import json
import re
import time
import threading
from pathlib import Path
from datetime import datetime
from collections import defaultdict

from PySide6.QtCore import Qt, QTimer, Signal, QObject, QRectF
from PySide6.QtGui import (
    QColor, QFont, QPainter, QPen, QPixmap, QFontDatabase, QBrush
)
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLabel, QPushButton,
    QVBoxLayout, QHBoxLayout, QGridLayout, QFrame, QMessageBox,
    QStackedWidget, QTableWidget, QTableWidgetItem, QHeaderView,
    QSpinBox, QCheckBox, QComboBox, QProgressBar,
    QDialog, QLineEdit, QFileDialog
)

# ============================================================
# SERIAL
# ============================================================

try:
    import serial
    import serial.tools.list_ports
except ImportError:
    serial = None


# ============================================================
# PATHS
# ============================================================

APP_DIR = Path(__file__).resolve().parent.parent
ASSETS = Path(__file__).resolve().parent / "assets"
DATA_DIR = APP_DIR / "data"

STATE_FILE = DATA_DIR / "counter_state.json"
SESSIONS_FILE = DATA_DIR / "work_sessions.csv"
LOG_FILE = DATA_DIR / "operations_log.csv"
CARTONS_FILE = DATA_DIR / "carton_events.csv"


# ============================================================
# CONFIG
# ============================================================

DEFAULT_TARGET = 4000
DEFAULT_WORK_MINUTES = 60

# EMA of the time per carton:
#   EMA_t = alpha * V_t + (1 - alpha) * EMA_(t-1),  alpha = 2 / (N + 1)
# N = cartons made in the look-back window at the nominal rate
# (1 carton / 6 s, 5 min -> 300 s / 6 = 50). alpha = 2 / 51 ~ 0.0392.
# The first N intervals are seeded with their simple average (SMA).
EMA_RATE_SECONDS = 6
EMA_WINDOW_MINUTES = 5
EMA_N = round(EMA_WINDOW_MINUTES * 60 / EMA_RATE_SECONDS)


# ============================================================
# IDENTITY
# ============================================================

CAMP_NAME = "Working Camp"
CAMP_NAME_AR = "معسكر العمل"

CAMP_TAGLINE = "Carton Counter"
CAMP_TAG_AR = "عدّاد الكراتين"

CAMP_YEAR_AR = "2026"

CAMP_MOTTO = "منه وله"

CAMP_VERSE = (
    "“For from Him and through Him and to Him are all things.”"
)
CAMP_REF = "Romans 11:36"

CAMP_VERSE_AR = "لأَنَّ مِنْهُ وَبِهِ وَلَهُ كُلَّ الأَشْيَاءِ"
CAMP_REF_AR = "رومية 11 : 36"


# ============================================================
# COLORS
# ============================================================

NAVY = "#14305F"
NAVY_DARK = "#0B2247"

OLIVE = "#3E4A24"
OLIVE_DARK = "#262F16"

KRAFT = "#DCC594"

PAPER = "#F3E6CC"
PAPER_2 = "#E9D8B4"
PAPER_3 = "#FBF5E6"

BROWN = "#3B2A1A"
BROWN_SOFT = "#7A5C3E"

GOLD = "#F2B824"

ORANGE = "#E48A17"

GREEN = "#2F8F46"

RED = "#B93A36"

BLUE = "#1F5FA8"

MAGENTA = "#D6336C"

GRAY = "#D9CDB4"

MUTED = "#7A6C58"

BORDER = "#CDB287"

FONT = "'Cairo','Tajawal','Segoe UI',Tahoma"


# ============================================================
# SERIAL CONFIGURATION
# ============================================================

INPUT_MODES = [
    (
        "count_total",
        "Running total: COUNT = 123 / 123"
    ),
    (
        "line_pulse",
        "Every received line = +1 carton"
    ),
    (
        "cts",
        "TTL pulse on CTS"
    ),
    (
        "dsr",
        "TTL pulse on DSR"
    ),
    (
        "dcd",
        "TTL pulse on DCD"
    ),
]

BAUD_CHOICES = [
    1200,
    2400,
    4800,
    9600,
    19200,
    38400,
    57600,
    115200
]

FIXED_PORT = "COM11"
FIXED_MODE = "cts"

DEFAULT_SERIAL = {
    "port": FIXED_PORT,
    "baud": 9600,
    "mode": FIXED_MODE,
    "debounce_ms": 0,
}

INPUT_MODES = [
    ("cts", "Laser beam on CTS — beam cut = +1 carton"),
]
# ============================================================
# HELPERS
# ============================================================

def qfont(size, weight=QFont.Bold, italic=False):
    f = QFont()

    f.setFamilies([
        "Cairo",
        "Tajawal",
        "Segoe UI",
        "Tahoma"
    ])

    f.setPixelSize(int(size))
    f.setWeight(weight)
    f.setItalic(italic)

    return f


def fmt_hms(seconds):
    seconds = max(0, int(seconds))

    return (
        f"{seconds // 3600:02d}:"
        f"{(seconds % 3600) // 60:02d}:"
        f"{seconds % 60:02d}"
    )


def ar_num(n):
    return f"{int(n):,}"


def ar_hms(seconds):
    return fmt_hms(seconds)


def load_pixmap(name, max_w, max_h):
    path = ASSETS / name

    if not path.exists():
        return None

    pix = QPixmap(str(path))

    if pix.isNull():
        return None

    return pix.scaled(
        max_w,
        max_h,
        Qt.KeepAspectRatio,
        Qt.SmoothTransformation
    )


def load_fonts():
    fonts = ASSETS / "fonts"

    if not fonts.exists():
        return

    for f in list(fonts.glob("*.ttf")) + list(fonts.glob("*.otf")):
        QFontDatabase.addApplicationFont(str(f))


# ============================================================
# COM PORT DISCOVERY
# ============================================================

def list_serial_ports():
    """
    Return:

        [
            ("COM11", "USB-SERIAL CH340"),
            ("COM5", "Silicon Labs CP210x"),
        ]

    """
    if serial is None:
        return []

    try:
        result = []

        for p in serial.tools.list_ports.comports():
            result.append(
                (
                    p.device,
                    p.description or p.device
                )
            )

        return result

    except Exception:
        return []


# ============================================================
# SERIAL BRIDGE
# ============================================================

class SerialBridge(QObject):

    delta_received = Signal(int)

    status_changed = Signal(str, bool)

    raw_data_received = Signal(str)

    def __init__(self, config=None):
        super().__init__()

        self.running = True

        self.serial_obj = None

        self.connected = False

        self.config = dict(DEFAULT_SERIAL)

        if config:
            self.config.update(config)

        self._lock_config()

        self._reconfigure = False

    def _lock_config(self):
        # Hardware is fixed: TTL-to-USB on COM11 sending 0/1 only.
        self.config["port"] = FIXED_PORT
        self.config["mode"] = FIXED_MODE

    # --------------------------------------------------------
    # CONFIG
    # --------------------------------------------------------

    def set_config(self, config):
        self.config = dict(DEFAULT_SERIAL, **config)

        self._lock_config()

        self._reconfigure = True

    # --------------------------------------------------------
    # STOP
    # --------------------------------------------------------

    def stop(self):
        self.running = False
        self._close()

    # --------------------------------------------------------
    # CLOSE
    # --------------------------------------------------------

    def _close(self):
        self.connected = False

        try:
            if self.serial_obj and self.serial_obj.is_open:
                self.serial_obj.close()
        except Exception:
            pass

        self.serial_obj = None

    # --------------------------------------------------------
    # PORT SCORING
    # --------------------------------------------------------

    def _ports(self):

        if serial is None:
            return []

        selected_port = self.config.get("port", "auto")

        if selected_port and selected_port != "auto":
            return [selected_port]

        try:
            ports = list(
                serial.tools.list_ports.comports()
            )
        except Exception:
            return []

        def score(p):

            text = (
                f"{p.description or ''} "
                f"{p.manufacturer or ''} "
                f"{p.hwid or ''}"
            ).lower()

            keys = {
                "ch340": 100,
                "ch341": 100,
                "wch": 90,
                "cp210": 90,
                "silicon labs": 90,
                "ftdi": 90,
                "ft232": 90,
                "pl2303": 90,
                "prolific": 80,
                "usb serial": 70,
                "usb-serial": 70,
                "uart": 60,
                "arduino": 50,
            }

            if "bluetooth" in text:
                return -100

            return sum(
                weight
                for key, weight in keys.items()
                if key in text
            )

        ports_sorted = sorted(
            ports,
            key=score,
            reverse=True
        )

        return [
            p.device
            for p in ports_sorted
            if score(p) > -100
        ]

    # --------------------------------------------------------
    # CONNECT
    # --------------------------------------------------------

    def _connect(self):

        if serial is None:

            self.status_changed.emit(
                "PySerial is not installed",
                False
            )

            return False

        baud = int(
            self.config.get(
                "baud",
                DEFAULT_SERIAL["baud"]
            )
        )

        devices = self._ports()
        print(f"[SERIAL PORTS] mode={self.config.get('mode', 'count_total')} baud={baud} candidates={devices!r}", flush=True)

        if not devices:

            self.status_changed.emit(
                "No COM port detected",
                False
            )

            return False

        for dev in devices:

            if not self.running:
                return False

            try:

                self.status_changed.emit(
                    f"Connecting • {dev} @ {baud}",
                    False
                )

                self.serial_obj = serial.Serial(
                    port=dev,
                    baudrate=baud,
                    timeout=0.05,
                    write_timeout=0.5
                )

                print(f"[SERIAL OPENED] port={dev} baud={baud} is_open={self.serial_obj.is_open}", flush=True)

                # Drive DTR/RTS high to power the chip's internal
                # reference that the laser sensor's CTS line needs.
                self.serial_obj.dtr = True
                self.serial_obj.rts = True

                time.sleep(0.3)

                try:
                    self.serial_obj.reset_input_buffer()
                except Exception:
                    pass

                self.connected = True

                self.status_changed.emit(
                    f"Sensor connected • {dev} @ {baud}",
                    True
                )

                return True

            except Exception as e:

                print(f"[SERIAL CONNECTION ERROR] {dev}: {type(e).__name__}: {e}", flush=True)
                self._close()

                continue

        self.status_changed.emit(
            "COM port found but connection failed",
            False
        )

        return False

    # --------------------------------------------------------
    # HARDWARE PIN
    # --------------------------------------------------------

    def _pin_state(self, mode):

        if not self.serial_obj:
            return False

        if mode == "cts":
            # Inverted: CTS drops when the laser beam is cut.
            return not bool(self.serial_obj.cts)

        if mode == "dsr":
            return bool(self.serial_obj.dsr)

        if mode == "dcd":
            return bool(self.serial_obj.cd)

        return False

    # --------------------------------------------------------
    # RUN
    # --------------------------------------------------------

    def run(self):

        last_total = None

        last_pin = None

        last_edge = 0.0

        buffer = b""

        while self.running:

            # ------------------------------------------------
            # RECONFIGURE
            # ------------------------------------------------

            if self._reconfigure:

                self._reconfigure = False

                self._close()

                last_total = None
                last_pin = None
                buffer = b""

            # ------------------------------------------------
            # CONNECT
            # ------------------------------------------------

            if not self.connected:

                if not self._connect():

                    time.sleep(1.5)

                    continue

                last_total = None
                last_pin = None
                buffer = b""

            mode = self.config.get(
                "mode",
                "count_total"
            )

            # ------------------------------------------------
            # HARDWARE PULSE MODES
            # ------------------------------------------------

            if mode in ("cts", "dsr", "dcd"):

                try:

                    pin = self._pin_state(mode)

                    now = time.monotonic()

                    debounce = max(
                        0,
                        int(
                            self.config.get(
                                "debounce_ms",
                                80
                            )
                        )
                    ) / 1000.0

                    # Rising edge
                    if (
                        last_pin is not None
                        and pin
                        and not last_pin
                        and now - last_edge >= debounce
                    ):

                        last_edge = now

                        self.delta_received.emit(1)

                        self.raw_data_received.emit(
                            f"{mode.upper()} pulse"
                        )

                    last_pin = pin

                    time.sleep(0.002)

                    continue

                except Exception as e:

                    print(f"[SERIAL PIN ERROR] mode={mode}: {type(e).__name__}: {e}", flush=True)
                    self._close()

                    self.status_changed.emit(
                        "USB disconnected • reconnecting...",
                        False
                    )

                    time.sleep(1)

                    continue

            # ------------------------------------------------
            # SERIAL DATA
            # ------------------------------------------------

            try:

                chunk = self.serial_obj.read(256)

                if not chunk:
                    continue

                print(f"[SERIAL BYTES] len={len(chunk)} raw={chunk!r} hex={chunk.hex(' ')}", flush=True)

                buffer += chunk

                # Protect against malformed streams
                if len(buffer) > 4096:
                    print(f"[SERIAL BUFFER WARNING] Dropping {len(buffer)} bytes without a complete line", flush=True)
                    buffer = b""

                parts = re.split(
                    rb"[\r\n]+",
                    buffer
                )

                buffer = parts.pop()

                for raw in parts:

                    line = raw.decode(
                        errors="ignore"
                    ).strip()

                    if not line:
                        continue

                    self.raw_data_received.emit(line)

                    # ----------------------------------------
                    # ONE LINE = ONE CARTON
                    # ----------------------------------------

                    if mode == "line_pulse":

                        self.delta_received.emit(1)

                        continue

                    # ----------------------------------------
                    # RUNNING TOTAL
                    # ----------------------------------------

                    if mode == "count_total":

                        numbers = re.findall(
                            r"\d+",
                            line
                        )

                        if not numbers:
                            print(f"[SERIAL PARSE] No number found in {line!r}", flush=True)
                            continue

                        total = int(numbers[-1])

                        if last_total is None:

                            print(f"[SERIAL TOTAL] Baseline/reset total={total}", flush=True)
                            last_total = total

                            continue

                        if total > last_total:

                            delta = total - last_total

                            self.delta_received.emit(delta)

                        # Device reset
                        elif total < last_total:

                            # Re-sync.
                            print(f"[SERIAL TOTAL] Baseline/reset total={total}", flush=True)
                            last_total = total

                            continue

                        last_total = total

            except Exception as e:

                print(f"[SERIAL READ ERROR] {type(e).__name__}: {e}", flush=True)
                self._close()

                self.status_changed.emit(
                    "USB disconnected • reconnecting...",
                    False
                )

                time.sleep(1.0)


# ============================================================
# INFO CARD
# ============================================================

class InfoCard(QFrame):

    def __init__(
        self,
        title,
        value="0",
        subtitle=""
    ):

        super().__init__()

        self.setObjectName("card")

        lay = QVBoxLayout(self)

        lay.setContentsMargins(
            16,
            12,
            16,
            12
        )

        lay.setSpacing(4)

        ttl = QLabel(title)

        ttl.setObjectName("infoTitle")

        ttl.setAlignment(Qt.AlignCenter)

        self.value = QLabel(value)

        self.value.setObjectName("infoValue")

        self.value.setAlignment(Qt.AlignCenter)

        self.subtitle = QLabel(subtitle)

        self.subtitle.setObjectName("infoSubtitle")

        self.subtitle.setAlignment(Qt.AlignCenter)

        lay.addWidget(ttl)

        lay.addWidget(
            self.value,
            1
        )

        lay.addWidget(
            self.subtitle
        )


# ============================================================
# DONUT
# ============================================================

class DonutWidget(QWidget):

    def __init__(self):

        super().__init__()

        self.count = 0

        self.target = DEFAULT_TARGET

        self.setMinimumSize(
            180,
            180
        )

    def set_values(
        self,
        count,
        target
    ):

        self.count = max(
            0,
            int(count)
        )

        self.target = max(
            1,
            int(target)
        )

        self.update()

    def paintEvent(self, _):

        p = QPainter(self)

        p.setRenderHint(
            QPainter.Antialiasing
        )

        side = min(
            self.width(),
            self.height()
        ) - 20

        rect = QRectF(
            (self.width() - side) / 2,
            (self.height() - side) / 2,
            side,
            side
        )

        width = max(
            14,
            int(side * 0.09)
        )

        pen = QPen(
            QColor(GRAY),
            width
        )

        pen.setCapStyle(
            Qt.RoundCap
        )

        p.setPen(pen)

        p.drawArc(
            rect.adjusted(
                width / 2,
                width / 2,
                -width / 2,
                -width / 2
            ),
            90 * 16,
            -360 * 16
        )

        pct = min(
            1.0,
            self.count / self.target
        )

        pen = QPen(
            QColor(ORANGE),
            width
        )

        pen.setCapStyle(
            Qt.RoundCap
        )

        p.setPen(pen)

        p.drawArc(
            rect.adjusted(
                width / 2,
                width / 2,
                -width / 2,
                -width / 2
            ),
            90 * 16,
            int(-360 * 16 * pct)
        )

        p.setPen(
            QColor(NAVY)
        )

        p.setFont(
            qfont(
                max(
                    20,
                    int(side * 0.16)
                )
            )
        )

        p.drawText(
            rect,
            Qt.AlignCenter,
            f"{round(pct * 100)}%"
        )


# ============================================================
# PIE CHART (display screen)
# ============================================================

class PieChartWidget(QWidget):

    def __init__(self):

        super().__init__()

        self.done = 0

        self.goal = DEFAULT_TARGET

        self.setMinimumSize(
            160,
            160
        )

    def set_values(
        self,
        done,
        goal
    ):

        self.done = max(
            0,
            int(done)
        )

        self.goal = max(
            1,
            int(goal)
        )

        self.update()

    def paintEvent(self, _):

        p = QPainter(self)

        p.setRenderHint(
            QPainter.Antialiasing
        )

        side = min(
            self.width(),
            self.height()
        ) - 12

        rect = QRectF(
            (self.width() - side) / 2,
            (self.height() - side) / 2,
            side,
            side
        )

        pct = min(
            1.0,
            self.done / self.goal
        )

        border = QPen(
            QColor(BROWN_SOFT),
            max(2, side * 0.012)
        )

        # Remaining slice
        p.setPen(border)

        p.setBrush(
            QBrush(QColor(GRAY))
        )

        p.drawEllipse(rect)

        # Done slice, clockwise from 12 o'clock
        if pct > 0:

            p.setBrush(
                QBrush(QColor(GREEN))
            )

            if pct >= 1.0:

                p.drawEllipse(rect)

            else:

                p.drawPie(
                    rect,
                    90 * 16,
                    int(-360 * 16 * pct)
                )

        # Percentage badge in the centre
        badge = side * 0.42

        badge_rect = QRectF(
            rect.center().x() - badge / 2,
            rect.center().y() - badge / 2,
            badge,
            badge
        )

        p.setBrush(
            QBrush(QColor(PAPER_3))
        )

        p.drawEllipse(badge_rect)

        p.setPen(
            QColor(NAVY)
        )

        p.setFont(
            qfont(
                max(
                    16,
                    int(side * 0.13)
                )
            )
        )

        p.drawText(
            badge_rect,
            Qt.AlignCenter,
            f"{round(pct * 100)}%"
        )


# ============================================================
# BAR CHART
# ============================================================

class BarChartWidget(QWidget):

    def __init__(self):

        super().__init__()

        self.values = []

        self.labels = []

        self.setMinimumHeight(220)

    def set_data(
        self,
        labels,
        values
    ):

        self.labels = labels[-5:]

        self.values = values[-5:]

        self.update()

    def paintEvent(self, _):

        p = QPainter(self)

        p.setRenderHint(
            QPainter.Antialiasing
        )

        p.fillRect(
            self.rect(),
            QColor(PAPER_3)
        )

        ml, mr, mt, mb = (
            48,
            20,
            26,
            48
        )

        area = QRectF(
            ml,
            mt,
            self.width() - ml - mr,
            self.height() - mt - mb
        )

        p.setPen(
            QPen(
                QColor("#C7B99F"),
                1
            )
        )

        p.drawLine(
            int(area.left()),
            int(area.bottom()),
            int(area.right()),
            int(area.bottom())
        )

        if not self.values:

            p.setPen(
                QColor(MUTED)
            )

            p.setFont(
                qfont(
                    14,
                    QFont.Normal
                )
            )

            p.drawText(
                area,
                Qt.AlignCenter,
                "No data yet"
            )

            return

        max_v = max(
            max(self.values),
            1
        )

        gap = 18

        n = len(self.values)

        bar_w = max(
            24,
            (area.width() - gap * (n + 1)) / n
        )

        for i, v in enumerate(
            self.values
        ):

            x = (
                area.left()
                + gap
                + i * (bar_w + gap)
            )

            h = (
                v
                / max_v
                * (area.height() - 24)
            )

            y = area.bottom() - h

            p.setPen(
                Qt.NoPen
            )

            p.setBrush(
                QColor(GOLD)
            )

            p.drawRoundedRect(
                QRectF(
                    x,
                    y,
                    bar_w,
                    h
                ),
                4,
                4
            )

            p.setPen(
                QColor(NAVY)
            )

            p.setFont(
                qfont(13)
            )

            p.drawText(
                QRectF(
                    x - 12,
                    y - 24,
                    bar_w + 24,
                    20
                ),
                Qt.AlignCenter,
                f"{int(v):,}"
            )

            label = (
                self.labels[i]
                if i < len(self.labels)
                else ""
            )

            p.setFont(
                qfont(
                    12,
                    QFont.Normal
                )
            )

            p.drawText(
                QRectF(
                    x - 18,
                    area.bottom() + 5,
                    bar_w + 36,
                    40
                ),
                Qt.AlignCenter,
                label
            )


# ============================================================
# STAMP CARD
# ============================================================

class StampCard(QFrame):

    def paintEvent(self, _):

        p = QPainter(self)

        p.setRenderHint(
            QPainter.Antialiasing
        )

        r = QRectF(
            self.rect()
        ).adjusted(
            10,
            10,
            -14,
            -14
        )

        p.setPen(
            Qt.NoPen
        )

        p.setBrush(
            QColor(
                60,
                40,
                10,
                55
            )
        )

        p.drawRoundedRect(
            r.translated(6, 8),
            8,
            8
        )

        p.setBrush(
            QColor(PAPER_3)
        )

        p.drawRoundedRect(
            r,
            8,
            8
        )

        p.setBrush(
            QColor(KRAFT)
        )

        rad = 7.0

        step = 24.0

        x = r.left() + step / 2

        while x < r.right():

            p.drawEllipse(
                QRectF(
                    x - rad,
                    r.top() - rad,
                    rad * 2,
                    rad * 2
                )
            )

            p.drawEllipse(
                QRectF(
                    x - rad,
                    r.bottom() - rad,
                    rad * 2,
                    rad * 2
                )
            )

            x += step

        y = r.top() + step / 2

        while y < r.bottom():

            p.drawEllipse(
                QRectF(
                    r.left() - rad,
                    y - rad,
                    rad * 2,
                    rad * 2
                )
            )

            p.drawEllipse(
                QRectF(
                    r.right() - rad,
                    y - rad,
                    rad * 2,
                    rad * 2
                )
            )

            y += step

        pen = QPen(
            QColor(BROWN_SOFT),
            1.4
        )

        p.setPen(pen)

        p.setBrush(
            Qt.NoBrush
        )

        p.drawRoundedRect(
            r.adjusted(
                20,
                20,
                -20,
                -20
            ),
            4,
            4
        )


# ============================================================
# COUNTER WINDOW
# ============================================================

class CounterWindow(QMainWindow):

    def __init__(
        self,
        state
    ):

        super().__init__()

        self.state = state

        self._scale = 0.0

        self.setWindowTitle(
            f"{CAMP_NAME} — Display"
        )

        self.setMinimumSize(
            900,
            600
        )

        root = QWidget()

        root.setObjectName(
            "counterRoot"
        )

        root.setLayoutDirection(
            Qt.RightToLeft
        )

        self.setCentralWidget(root)

        outer = QVBoxLayout(root)

        outer.setContentsMargins(
            26,
            22,
            26,
            26
        )

        card = StampCard()

        outer.addWidget(card)

        lay = QVBoxLayout(card)

        lay.setContentsMargins(
            70,
            52,
            70,
            46
        )

        lay.setSpacing(6)

        header = QHBoxLayout()

        header.setSpacing(22)

        self.logo_label = self._make_logo(
            120,
            150
        )

        header.addWidget(
            self.logo_label
        )

        name_box = QVBoxLayout()

        name_box.setSpacing(0)

        self.name_ar = QLabel(
            CAMP_NAME_AR
        )

        self.name_ar.setObjectName(
            "cName"
        )

        self.tag_ar = QLabel(
            f"{CAMP_TAG_AR}  •  احتياج {CAMP_YEAR_AR}"
        )

        self.tag_ar.setObjectName(
            "cTag"
        )

        self.tag_en = QLabel(
            f"{CAMP_NAME.upper()}  ·  {CAMP_TAGLINE.upper()}"
        )

        self.tag_en.setObjectName(
            "cTagEn"
        )

        self.tag_en.setLayoutDirection(
            Qt.LeftToRight
        )

        self.tag_en.setAlignment(
            Qt.AlignRight | Qt.AlignVCenter
        )

        name_box.addWidget(
            self.name_ar
        )

        name_box.addWidget(
            self.tag_ar
        )

        name_box.addWidget(
            self.tag_en
        )

        header.addLayout(
            name_box
        )

        header.addStretch(1)

        self.status_label = QLabel(
            "الحساس غير متصل"
        )

        self.status_label.setObjectName(
            "cStatus"
        )

        header.addWidget(
            self.status_label,
            0,
            Qt.AlignTop | Qt.AlignLeft
        )

        lay.addLayout(header)

        lay.addStretch(2)

        self.caption = QLabel(
            "عدد الكراتين اللي جهزناها"
        )

        self.caption.setObjectName(
            "cCaption"
        )

        self.caption.setAlignment(
            Qt.AlignCenter
        )

        lay.addWidget(
            self.caption
        )

        band = QFrame()

        band.setObjectName(
            "cBand"
        )

        bl = QVBoxLayout(band)

        bl.setContentsMargins(
            30,
            0,
            30,
            0
        )

        self.count_label = QLabel(
            "0"
        )

        self.count_label.setObjectName(
            "cBig"
        )

        self.count_label.setAlignment(
            Qt.AlignCenter
        )

        self.count_label.setTextFormat(
            Qt.RichText
        )

        bl.addWidget(
            self.count_label
        )

        # Count band + pie chart side by side (RTL: band on the right)
        mid_row = QHBoxLayout()

        mid_row.setSpacing(40)

        mid_row.addStretch(1)

        mid_row.addWidget(
            band,
            0,
            Qt.AlignVCenter
        )

        pie_box = QVBoxLayout()

        pie_box.setSpacing(4)

        self.pie_title = QLabel()

        self.pie_title.setObjectName(
            "cPieTitle"
        )

        self.pie_title.setAlignment(
            Qt.AlignCenter
        )

        self.pie = PieChartWidget()

        self.pie_legend = QLabel()

        self.pie_legend.setObjectName(
            "cPieLegend"
        )

        self.pie_legend.setAlignment(
            Qt.AlignCenter
        )

        pie_box.addWidget(
            self.pie_title
        )

        pie_box.addWidget(
            self.pie,
            0,
            Qt.AlignCenter
        )

        pie_box.addWidget(
            self.pie_legend
        )

        mid_row.addLayout(
            pie_box
        )

        mid_row.addStretch(1)

        lay.addLayout(
            mid_row
        )

        lay.addSpacing(10)

        self.goal_label = QLabel()

        self.goal_label.setObjectName(
            "cGoal"
        )

        self.goal_label.setAlignment(
            Qt.AlignCenter
        )

        lay.addWidget(
            self.goal_label
        )

        prog_row = QHBoxLayout()

        prog_row.setSpacing(18)

        self.bar = QProgressBar()

        self.bar.setObjectName(
            "cBar"
        )

        self.bar.setRange(
            0,
            1000
        )

        self.bar.setTextVisible(
            False
        )

        self.progress_label = QLabel(
            "0%"
        )

        self.progress_label.setObjectName(
            "cPct"
        )

        prog_row.addWidget(
            self.bar,
            1
        )

        prog_row.addWidget(
            self.progress_label
        )

        lay.addLayout(
            prog_row
        )

        lay.addSpacing(4)

        countdown_row = QHBoxLayout()

        countdown_row.setSpacing(14)

        self.countdown_caption = QLabel(
            "الوقت المتبقي لفترة الشغل"
        )

        self.countdown_caption.setObjectName(
            "cCdCaption"
        )

        self.countdown_value = QLabel(
            "00:00:00"
        )

        self.countdown_value.setObjectName(
            "cCdValue"
        )

        countdown_row.addStretch(1)

        countdown_row.addWidget(
            self.countdown_caption
        )

        countdown_row.addWidget(
            self.countdown_value
        )

        countdown_row.addStretch(1)

        lay.addLayout(
            countdown_row
        )

        # Time per carton: EMA (now) • plain average • ideal (required)
        rates_row = QHBoxLayout()

        rates_row.setSpacing(14)

        rates_row.addStretch(1)

        def rate_pair(caption):

            cap = QLabel(caption)

            cap.setObjectName(
                "cCdCaption"
            )

            val = QLabel("0.00 ث")

            val.setObjectName(
                "cCdValue"
            )

            rates_row.addWidget(cap)

            rates_row.addWidget(val)

            return val

        self.rate_value = rate_pair(
            "زمن الكرتونة دلوقتي (EMA)"
        )

        rates_row.addSpacing(40)

        self.avg_value = rate_pair(
            "المتوسط العادي"
        )

        rates_row.addSpacing(40)

        self.ideal_value = rate_pair(
            "المطلوب (Ideal)"
        )

        rates_row.addStretch(1)

        lay.addLayout(
            rates_row
        )

        lay.addStretch(2)

        footer = QHBoxLayout()

        footer.setSpacing(24)

        verse_box = QVBoxLayout()

        verse_box.setSpacing(0)

        v = QLabel(
            CAMP_VERSE_AR
        )

        v.setObjectName(
            "cVerse"
        )

        v.setWordWrap(True)

        r = QLabel(
            f"({CAMP_REF_AR})"
        )

        r.setObjectName(
            "cRef"
        )

        en = QLabel(
            f"{CAMP_VERSE}  ({CAMP_REF})"
        )

        en.setObjectName(
            "cVerseEn"
        )

        en.setLayoutDirection(
            Qt.LeftToRight
        )

        en.setWordWrap(True)

        verse_box.addWidget(v)
        verse_box.addWidget(r)
        verse_box.addWidget(en)

        footer.addLayout(
            verse_box,
            1
        )

        self.motto = QLabel(
            CAMP_MOTTO
        )

        self.motto.setObjectName(
            "cMotto"
        )

        self.motto.setAlignment(
            Qt.AlignCenter
        )

        footer.addWidget(
            self.motto,
            0,
            Qt.AlignBottom
        )

        lay.addLayout(
            footer
        )

        self.apply_styles(1.0)

        self.refresh()

    def _make_logo(
        self,
        w,
        h
    ):

        lbl = QLabel()

        lbl.setAlignment(
            Qt.AlignCenter
        )

        pix = load_pixmap(
            "mark.png",
            w,
            h
        )

        if pix:

            lbl.setPixmap(pix)

        else:

            lbl.setText(
                CAMP_NAME_AR
            )

            lbl.setObjectName(
                "cLogoFallback"
            )

        return lbl

    def apply_styles(
        self,
        k
    ):

        s = lambda v: int(v * k)

        self.setStyleSheet(
            f"""
            QWidget#counterRoot {{
                background: {KRAFT};
            }}

            QLabel {{
                background: transparent;
                font-family: {FONT};
            }}

            QLabel#cName {{
                color: {NAVY};
                font-size: {s(64)}px;
                font-weight: 900;
            }}

            QLabel#cTag {{
                color: {OLIVE};
                font-size: {s(24)}px;
                font-weight: 800;
            }}

            QLabel#cTagEn {{
                color: {BROWN_SOFT};
                font-size: {s(14)}px;
                font-weight: 700;
                letter-spacing: 3px;
            }}

            QLabel#cLogoFallback {{
                color: {NAVY};
                font-size: {s(26)}px;
                font-weight: 900;
                border: 2px dashed {BROWN_SOFT};
                border-radius: 10px;
                padding: 16px;
            }}

            QLabel#cStatus {{
                color: {RED};
                font-size: {s(18)}px;
                font-weight: 800;
            }}

            QLabel#cCaption {{
                color: {OLIVE};
                font-size: {s(34)}px;
                font-weight: 800;
            }}

            QFrame#cBand {{
                background: {GREEN};
                border-radius: {s(30)}px;
            }}

            QLabel#cBig {{
                color: {PAPER_3};
                font-size: {s(220)}px;
                font-weight: 900;
            }}

            QLabel#cGoal {{
                color: {BROWN};
                font-size: {s(30)}px;
                font-weight: 800;
            }}

            QLabel#cPieTitle {{
                color: {OLIVE};
                font-size: {s(24)}px;
                font-weight: 900;
            }}

            QLabel#cPieLegend {{
                color: {BROWN};
                font-size: {s(20)}px;
                font-weight: 800;
            }}

            QProgressBar#cBar {{
                background: {GRAY};
                border: 2px solid {BROWN_SOFT};
                border-radius: {s(14)}px;
                min-height: {s(28)}px;
                max-height: {s(28)}px;
            }}

            QProgressBar#cBar::chunk {{
                background: {ORANGE};
                border-radius: {s(10)}px;
                margin: 2px;
            }}

            QLabel#cPct {{
                color: {NAVY};
                font-size: {s(44)}px;
                font-weight: 900;
            }}

            QLabel#cCdCaption {{
                color: {OLIVE};
                font-size: {s(20)}px;
                font-weight: 800;
            }}

            QLabel#cCdValue {{
                color: {NAVY};
                font-size: {s(30)}px;
                font-weight: 900;
            }}

            QLabel#cCdValue[over="true"] {{
                color: {RED};
            }}

            QLabel#cVerse {{
                color: {NAVY};
                font-size: {s(26)}px;
                font-weight: 800;
            }}

            QLabel#cRef {{
                color: {ORANGE};
                font-size: {s(18)}px;
                font-weight: 800;
            }}

            QLabel#cVerseEn {{
                color: {BROWN_SOFT};
                font-size: {s(13)}px;
                font-style: italic;
                font-weight: 600;
            }}

            QLabel#cMotto {{
                color: {BROWN};
                font-size: {s(40)}px;
                font-weight: 900;
                border: 4px solid {BROWN};
                border-radius: 6px;
                padding: {s(6)}px {s(22)}px;
            }}
            """
        )

        self.pie.setFixedSize(
            s(330),
            s(330)
        )

    def resizeEvent(self, e):

        k = max(
            0.6,
            min(
                1.6,
                self.height() / 1080
            )
        )

        if abs(k - self._scale) > 0.03:

            self._scale = k

            self.apply_styles(k)

            self.refresh()

        super().resizeEvent(e)

    def refresh(self):

        count = self.state["count"]

        k = self._scale or 1.0

        self.count_label.setText(
            f"{ar_num(count)} "
            f"<span style='font-size:{int(72 * k)}px;'>كرتونة</span>"
        )

        # The display shows the session goal only; the full goal
        # stays on the dashboard.
        session = self.state.get("active_session")

        if not session:

            self.bar.setValue(0)

            self.progress_label.setText(
                "0%"
            )

            self.goal_label.setText(
                "مفيش جلسة شغالة دلوقتي"
            )

            self.pie_title.hide()

            self.pie.hide()

            self.pie_legend.hide()

            return

        done = max(
            0,
            count - int(session.get("start_count", 0))
        )

        goal = max(
            1,
            int(session.get("goal", 1))
        )

        pct = min(
            1.0,
            done / goal
        )

        name = session.get("name") or "الجلسة"

        self.bar.setValue(
            int(pct * 1000)
        )

        self.progress_label.setText(
            f"{round(pct * 100)}%"
        )

        if done >= goal:

            self.goal_label.setText(
                f"تم الوصول لهدف {name} "
                f"({ar_num(goal)} كرتونة)"
                f" — {CAMP_MOTTO}"
            )

        else:

            self.goal_label.setText(
                f"هدف {name} {ar_num(goal)} كرتونة"
                f" • "
                f"المتبقي {ar_num(goal - done)}"
            )

        self.pie_title.setText(
            f"هدف {name}"
        )

        self.pie.set_values(
            done,
            goal
        )

        self.pie_legend.setText(
            f"تم {ar_num(done)}"
            f" • "
            f"متبقي {ar_num(max(0, goal - done))}"
        )

        show_pie = bool(
            self.state.get("show_pie", True)
        )

        self.pie_title.setVisible(show_pie)

        self.pie.setVisible(show_pie)

        self.pie_legend.setVisible(show_pie)

    def set_status(
        self,
        text,
        connected
    ):

        if connected:

            self.status_label.setText(
                "الحساس متصل"
            )

        else:

            self.status_label.setText(
                "الحساس غير متصل"
            )

        color = (
            GREEN
            if connected
            else RED
        )

        self.status_label.setStyleSheet(
            f"""
            color: {color};
            font-family: {FONT};
            font-size: {int(18 * (self._scale or 1.0))}px;
            font-weight: 800;
            """
        )

    def set_rate(
        self,
        ema,
        avg,
        ideal,
        cartons
    ):

        self.rate_value.setText(
            f"{ema:.2f} ث"
            if cartons > 0
            else "0.00 ث"
        )

        self.avg_value.setText(
            f"{avg:.2f} ث"
        )

        self.ideal_value.setText(
            f"{ideal:.2f} ث"
        )

        # EMA slower than ideal = behind pace -> red
        behind = (
            cartons > 0
            and ideal > 0
            and ema > ideal
        )

        self.rate_value.setProperty(
            "over",
            behind
        )

        self.rate_value.style().unpolish(
            self.rate_value
        )

        self.rate_value.style().polish(
            self.rate_value
        )

    def set_countdown(
        self,
        seconds_remaining,
        over,
        active
    ):

        if not active:

            self.countdown_value.setText(
                "00:00:00"
            )

            self.countdown_caption.setText(
                "الوقت المتبقي لفترة الشغل"
                " • الفترة لسه مبدأتش"
            )

        elif over:

            self.countdown_value.setText(
                "00:00:00"
            )

            self.countdown_caption.setText(
                "انتهت فترة الشغل"
            )

        else:

            self.countdown_value.setText(
                ar_hms(seconds_remaining)
            )

            self.countdown_caption.setText(
                "الوقت المتبقي لفترة الشغل"
            )

        self.countdown_value.setProperty(
            "over",
            bool(
                active and over
            )
        )

        self.countdown_value.style().unpolish(
            self.countdown_value
        )

        self.countdown_value.style().polish(
            self.countdown_value
        )


# ============================================================
# EXCEL EXPORT
# ============================================================

def read_carton_events():

    if not CARTONS_FILE.exists():
        return []

    try:

        with CARTONS_FILE.open(
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as f:

            return list(
                csv.DictReader(f)
            )

    except Exception:

        return []


def export_workbook(
    path,
    sessions,
    events,
    full_target
):
    """
    Summary sheet: one row per session with its goal, its weight
    of the full goal, and its EMA time/carton.
    Cycle Times sheet: every timed carton with the EMA as live
    Excel formulas driven by N on the Summary sheet.
    """

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    def parse_dt(text):
        try:
            return datetime.fromisoformat(text)
        except Exception:
            return text or None

    def num(text, default=0.0):
        try:
            return float(text)
        except Exception:
            return default

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3E4A24")
    bold = Font(bold=True)

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

    wb = Workbook()

    # ---------------- Cycle Times (built first: Summary links to it)
    cy = wb.active
    cy.title = "Cycle Times"

    cy_headers = [
        "Session",
        "Session Start",
        "Carton #",
        "Time",
        "Cycle Time (s)",
        "EMA (s)",
        "Running Avg (s)",
        "Ideal (s)",
    ]

    cy.append(cy_headers)
    style_header(cy, 1, len(cy_headers))

    by_session = defaultdict(list)

    for e in events:
        by_session[e.get("session_id", "")].append(e)

    last_ema_cell = {}

    # Ideal pace per session = planned time / goal
    ideal_by_session = {}

    for s_row in sessions:

        g = num(s_row.get("goal"), 0)

        m = num(s_row.get("planned_minutes"), 0)

        if g > 0 and m > 0:

            ideal_by_session[
                s_row.get("session_id") or s_row.get("start_time", "")
            ] = m * 60 / g

    r = 2

    for start_key, rows in by_session.items():

        first_cycle_row = None

        for e in rows:

            cycle = e.get("cycle_seconds", "")

            cy.cell(r, 1, e.get("session_name") or "-")
            cy.cell(r, 2, parse_dt(start_key)).number_format = "yyyy-mm-dd hh:mm"
            cy.cell(r, 3, int(num(e.get("carton"), 0)))
            cy.cell(r, 4, parse_dt(e.get("time", ""))).number_format = "hh:mm:ss"

            if cycle != "":

                cy.cell(r, 5, num(cycle)).number_format = "0.00"

                if first_cycle_row is None:
                    first_cycle_row = r

                # SMA for the first N cycles, then the recursive EMA
                cy.cell(
                    r,
                    6,
                    f"=IF(COUNT(E${first_cycle_row}:E{r})<=Summary!$B$3,"
                    f"AVERAGE(E${first_cycle_row}:E{r}),"
                    f"Summary!$B$4*E{r}+(1-Summary!$B$4)*F{r - 1})"
                ).number_format = "0.00"

                last_ema_cell[start_key] = f"'Cycle Times'!F{r}"

                cy.cell(
                    r,
                    7,
                    f"=AVERAGE(E${first_cycle_row}:E{r})"
                ).number_format = "0.00"

                if start_key in ideal_by_session:

                    cy.cell(
                        r,
                        8,
                        ideal_by_session[start_key]
                    ).number_format = "0.00"

            r += 1

    for col, width in zip("ABCDEFGH", [18, 18, 10, 12, 15, 12, 15, 12]):
        cy.column_dimensions[col].width = width

    cy.freeze_panes = "A2"

    # ---------------- Summary
    sm = wb.create_sheet("Summary", 0)

    sm["A1"] = f"{CAMP_NAME} — Sessions Report"
    sm["A1"].font = Font(bold=True, size=14)

    # N = cartons in the look-back window at the nominal rate
    sm["D3"] = "Nominal rate (s / carton)"
    sm["F3"] = EMA_RATE_SECONDS

    sm["D4"] = "Recovery window (min)"
    sm["F4"] = EMA_WINDOW_MINUTES

    for a in ("D3", "D4"):
        sm[a].font = bold

    sm["A3"] = "EMA window N (cartons)"
    sm["B3"] = "=ROUND(F4*60/F3,0)"

    sm["A4"] = "Smoothing factor α = 2/(N+1)"
    sm["B4"] = "=2/(B3+1)"
    sm["B4"].number_format = "0.0000"

    sm["A5"] = "Full goal (cartons)"
    sm["B5"] = int(full_target)

    sm["A7"] = "Ideal"
    sm["B7"] = (
        "Ideal time/carton = planned session time (s) ÷ session goal"
    )
    sm["A7"].font = Font(bold=True)

    sm["A6"] = "EMA formula"
    sm["B6"] = (
        "EMA_t = α × V_t + (1 − α) × EMA_(t−1)   "
        "(first N cycles seeded with their simple average)"
    )

    for a in ("A3", "A4", "A5", "A6"):
        sm[a].font = bold

    hdr_row = 8

    sm_headers = [
        "Session",
        "Start",
        "End",
        "Duration (min)",
        "Cartons",
        "Session Goal",
        "Goal Reached %",
        "Weight of Full Goal",
        "Contribution to Full Goal",
        "Avg Time / Carton (s)",
        "EMA Time / Carton (s)",
        "Planned Time (min)",
        "Ideal Time / Carton (s)",
        "EMA − Ideal (s)",
    ]

    for c, h in enumerate(sm_headers, 1):
        sm.cell(hdr_row, c, h)

    style_header(sm, hdr_row, len(sm_headers))

    r = hdr_row + 1

    for s in sessions:

        start_key = s.get("start_time", "")

        session_key = s.get("session_id") or start_key

        goal = int(num(s.get("goal"), 0))

        sm.cell(r, 1, s.get("name") or "-")
        sm.cell(r, 2, parse_dt(start_key)).number_format = "yyyy-mm-dd hh:mm"
        sm.cell(r, 3, parse_dt(s.get("end_time", ""))).number_format = "yyyy-mm-dd hh:mm"
        sm.cell(r, 4, num(s.get("duration_seconds")) / 60).number_format = "0.0"
        sm.cell(r, 5, int(num(s.get("cartons"), 0)))
        sm.cell(r, 6, goal if goal else None)

        sm.cell(r, 7, f'=IF(N(F{r})>0,E{r}/F{r},"")').number_format = "0.0%"
        sm.cell(r, 8, f'=IF(N(F{r})>0,F{r}/$B$5,"")').number_format = "0.0%"
        sm.cell(r, 9, f"=E{r}/$B$5").number_format = "0.0%"
        sm.cell(r, 10, f'=IF(E{r}>0,D{r}*60/E{r},"")').number_format = "0.00"

        if session_key in last_ema_cell:

            sm.cell(r, 11, f"={last_ema_cell[session_key]}")

        else:

            ema = num(s.get("ema_seconds_per_carton"), 0)

            sm.cell(r, 11, ema if ema else None)

        sm.cell(r, 11).number_format = "0.00"

        planned = num(s.get("planned_minutes"), 0)

        sm.cell(r, 12, planned if planned else None)

        # Ideal = planned seconds / goal; positive gap = slower than ideal
        sm.cell(
            r,
            13,
            f'=IF(AND(N(L{r})>0,N(F{r})>0),L{r}*60/F{r},"")'
        ).number_format = "0.00"

        sm.cell(
            r,
            14,
            f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(M{r})),K{r}-M{r},"")'
        ).number_format = "+0.00;-0.00;0.00"

        r += 1

    # Totals
    if r > hdr_row + 1:

        first, last = hdr_row + 1, r - 1

        sm.cell(r, 1, "Total").font = bold
        sm.cell(r, 5, f"=SUM(E{first}:E{last})").font = bold
        sm.cell(r, 6, f"=SUM(F{first}:F{last})").font = bold

        c = sm.cell(r, 8, f"=SUM(H{first}:H{last})")
        c.font = bold
        c.number_format = "0.0%"

        c = sm.cell(r, 9, f"=SUM(I{first}:I{last})")
        c.font = bold
        c.number_format = "0.0%"

    widths = [18, 17, 17, 14, 10, 13, 14, 16, 18, 16, 16, 14, 16, 14]

    for i, w in enumerate(widths, 1):
        sm.column_dimensions[get_column_letter(i)].width = w

    sm.column_dimensions["A"].width = 30

    sm.row_dimensions[hdr_row].height = 32

    sm.freeze_panes = sm.cell(hdr_row + 1, 1)

    wb.save(path)


# ============================================================
# DIALOGS
# ============================================================

def _dialog_buttons(dlg, ok_text):

    row = QHBoxLayout()

    row.addStretch(1)

    cancel = QPushButton("Cancel")

    cancel.setObjectName(
        "secondaryButton"
    )

    ok = QPushButton(ok_text)

    ok.setObjectName(
        "primaryButton"
    )

    ok.setDefault(True)

    cancel.clicked.connect(dlg.reject)

    ok.clicked.connect(dlg.accept)

    row.addWidget(cancel)

    row.addWidget(ok)

    return row


def _field_label(text):

    lbl = QLabel(text)

    lbl.setObjectName(
        "settingsTitle"
    )

    return lbl


class SessionDialog(QDialog):

    def __init__(
        self,
        parent,
        name="",
        goal=500,
        minutes=DEFAULT_WORK_MINUTES
    ):

        super().__init__(parent)

        self.setWindowTitle(
            "New Session"
        )

        self.setMinimumWidth(
            420
        )

        lay = QVBoxLayout(self)

        lay.setContentsMargins(
            22,
            18,
            22,
            18
        )

        lay.setSpacing(8)

        lay.addWidget(
            _field_label("Session name (optional)")
        )

        self.name_edit = QLineEdit(name)

        self.name_edit.setPlaceholderText(
            "e.g. Morning shift"
        )

        lay.addWidget(
            self.name_edit
        )

        lay.addWidget(
            _field_label("Goal (cartons)")
        )

        self.goal_spin = QSpinBox()

        self.goal_spin.setRange(
            1,
            1_000_000
        )

        self.goal_spin.setValue(
            max(1, goal)
        )

        lay.addWidget(
            self.goal_spin
        )

        lay.addWidget(
            _field_label("Time (minutes)")
        )

        self.minutes_spin = QSpinBox()

        self.minutes_spin.setRange(
            1,
            24 * 60
        )

        self.minutes_spin.setSuffix(
            " min"
        )

        self.minutes_spin.setValue(
            max(1, minutes)
        )

        lay.addWidget(
            self.minutes_spin
        )

        lay.addSpacing(10)

        lay.addLayout(
            _dialog_buttons(
                self,
                "Start Session"
            )
        )

    def values(self):

        return (
            self.name_edit.text().strip(),
            int(self.goal_spin.value()),
            int(self.minutes_spin.value()),
        )


class ManualEntryDialog(QDialog):

    def __init__(
        self,
        parent,
        current
    ):

        super().__init__(parent)

        self.setWindowTitle(
            "Manual Entry"
        )

        self.setMinimumWidth(
            400
        )

        lay = QVBoxLayout(self)

        lay.setContentsMargins(
            22,
            18,
            22,
            18
        )

        lay.setSpacing(8)

        info = QLabel(
            f"Current total: {current:,} cartons"
        )

        info.setObjectName(
            "settingsDesc"
        )

        lay.addWidget(info)

        lay.addWidget(
            _field_label("Action")
        )

        self.mode_combo = QComboBox()

        self.mode_combo.addItem(
            "Add cartons",
            "add"
        )

        self.mode_combo.addItem(
            "Subtract cartons",
            "subtract"
        )

        self.mode_combo.addItem(
            "Set total to",
            "set"
        )

        lay.addWidget(
            self.mode_combo
        )

        lay.addWidget(
            _field_label("Amount")
        )

        self.amount_spin = QSpinBox()

        self.amount_spin.setRange(
            0,
            1_000_000
        )

        self.amount_spin.setValue(0)

        lay.addWidget(
            self.amount_spin
        )

        lay.addSpacing(10)

        lay.addLayout(
            _dialog_buttons(
                self,
                "Apply"
            )
        )

        self.amount_spin.setFocus()

        self.amount_spin.selectAll()

    def values(self):

        return (
            self.mode_combo.currentData(),
            int(self.amount_spin.value()),
        )


# ============================================================
# DASHBOARD
# ============================================================

class DashboardWindow(QMainWindow):

    count_changed = Signal(int)

    countdown_changed = Signal(
        int,
        bool,
        bool
    )

    # ema, plain average, ideal (seconds per carton), timed cartons
    rate_changed = Signal(
        float,
        float,
        float,
        int
    )

    serial_config_changed = Signal(dict)

    def __init__(
        self,
        state
    ):

        super().__init__()

        self.state = state

        self.setWindowTitle(
            f"{CAMP_NAME_AR} — {CAMP_NAME} — Dashboard"
        )

        self.setMinimumSize(
            1200,
            760
        )

        self.session_active = False

        self.session_start_count = 0

        self.session_start_mono = None

        self.session_start_wall = None

        self.session_name = ""

        self.session_goal = 0

        self.session_minutes = DEFAULT_WORK_MINUTES

        self._reset_ema()

        self.sensor_connected = False

        root = QWidget()

        root.setObjectName(
            "page"
        )

        self.setCentralWidget(root)

        outer = QHBoxLayout(root)

        outer.setContentsMargins(
            0,
            0,
            0,
            0
        )

        outer.setSpacing(0)

        outer.addWidget(
            self.build_sidebar()
        )

        content = QWidget()

        cl = QVBoxLayout(content)

        cl.setContentsMargins(
            0,
            0,
            0,
            0
        )

        cl.setSpacing(0)

        cl.addWidget(
            self.build_header()
        )

        self.pages = QStackedWidget()

        cl.addWidget(
            self.pages,
            1
        )

        outer.addWidget(
            content,
            1
        )

        self.home_page = self.build_home()

        self.reports_page = self.build_reports()

        self.log_page = self.build_log()

        self.settings_page = self.build_settings()

        for p in [
            self.home_page,
            self.reports_page,
            self.log_page,
            self.settings_page
        ]:

            self.pages.addWidget(p)

        self.apply_styles()

        self.connect_actions()

        self.refresh_all()

        self.timer = QTimer(self)

        self.timer.timeout.connect(
            self.tick
        )

        self.timer.start(500)

    # ========================================================
    # SIDEBAR
    # ========================================================

    def make_nav(self, text):

        btn = QPushButton(text)

        btn.setObjectName(
            "navButton"
        )

        btn.setCheckable(True)

        btn.setCursor(
            Qt.PointingHandCursor
        )

        btn.setMinimumHeight(
            52
        )

        return btn

    def build_sidebar(self):

        side = QFrame()

        side.setObjectName(
            "sidebar"
        )

        side.setFixedWidth(
            210
        )

        lay = QVBoxLayout(side)

        lay.setContentsMargins(
            10,
            16,
            10,
            16
        )

        lay.setSpacing(6)

        self.nav_home = self.make_nav(
            "Home"
        )

        self.nav_reports = self.make_nav(
            "Reports"
        )

        self.nav_log = self.make_nav(
            "Operation Log"
        )

        self.nav_settings = self.make_nav(
            "Settings"
        )

        self.nav_exit = self.make_nav(
            "Exit"
        )

        self.nav_home.setChecked(
            True
        )

        for b in [
            self.nav_home,
            self.nav_reports,
            self.nav_log,
            self.nav_settings,
            self.nav_exit
        ]:

            lay.addWidget(b)

        lay.addStretch(1)

        footer = QFrame()

        footer.setObjectName(
            "sideFooter"
        )

        fl = QVBoxLayout(footer)

        fl.setContentsMargins(
            10,
            12,
            10,
            10
        )

        logo = QLabel()

        logo.setAlignment(
            Qt.AlignCenter
        )

        pix = load_pixmap(
            "mark.png",
            90,
            110
        )

        if pix:

            logo.setPixmap(pix)

        title = QLabel(
            CAMP_NAME_AR
        )

        title.setObjectName(
            "sideTitle"
        )

        title.setAlignment(
            Qt.AlignCenter
        )

        sub = QLabel(
            f"{CAMP_NAME} • {CAMP_TAGLINE}"
        )

        sub.setObjectName(
            "sideSub"
        )

        sub.setAlignment(
            Qt.AlignCenter
        )

        sub.setWordWrap(True)

        self.usb_label = QLabel(
            "Sensor not connected"
        )

        self.usb_label.setObjectName(
            "usbStatus"
        )

        self.usb_label.setAlignment(
            Qt.AlignCenter
        )

        self.usb_label.setWordWrap(True)

        fl.addWidget(logo)

        fl.addWidget(title)

        fl.addWidget(sub)

        fl.addSpacing(6)

        fl.addWidget(
            self.usb_label
        )

        lay.addWidget(
            footer
        )

        return side

    # ========================================================
    # HEADER
    # ========================================================

    def build_header(self):

        header = QFrame()

        header.setObjectName(
            "header"
        )

        header.setMinimumHeight(
            96
        )

        lay = QHBoxLayout(header)

        lay.setContentsMargins(
            20,
            8,
            20,
            8
        )

        lay.setSpacing(16)

        logo = QLabel()

        logo.setAlignment(
            Qt.AlignCenter
        )

        pix = load_pixmap(
            "mark.png",
            64,
            80
        )

        if pix:

            logo.setPixmap(pix)

        names = QVBoxLayout()

        names.setSpacing(0)

        ar = QLabel(
            CAMP_NAME_AR
        )

        ar.setObjectName(
            "headerTitle"
        )

        en = QLabel(
            f"{CAMP_NAME.upper()} • {CAMP_TAGLINE.upper()}"
        )

        en.setObjectName(
            "headerSub"
        )

        names.addStretch(1)

        names.addWidget(ar)

        names.addWidget(en)

        names.addStretch(1)

        verse = QLabel(
            f"{CAMP_VERSE} ({CAMP_REF})"
        )

        verse.setObjectName(
            "verse"
        )

        verse.setAlignment(
            Qt.AlignCenter
        )

        verse.setWordWrap(True)

        motto = QLabel(
            CAMP_MOTTO
        )

        motto.setObjectName(
            "motto"
        )

        motto.setAlignment(
            Qt.AlignCenter
        )

        time_box = QFrame()

        time_box.setObjectName(
            "timeBox"
        )

        time_box.setFixedWidth(
            170
        )

        tl = QVBoxLayout(time_box)

        tl.setContentsMargins(
            10,
            6,
            10,
            6
        )

        tl.setSpacing(0)

        self.date_label = QLabel()

        self.date_label.setObjectName(
            "date"
        )

        self.date_label.setAlignment(
            Qt.AlignCenter
        )

        self.time_label = QLabel()

        self.time_label.setObjectName(
            "clock"
        )

        self.time_label.setAlignment(
            Qt.AlignCenter
        )

        tl.addWidget(
            self.date_label
        )

        tl.addWidget(
            self.time_label
        )

        lay.addWidget(logo)

        lay.addLayout(names)

        lay.addWidget(
            verse,
            1
        )

        lay.addWidget(motto)

        lay.addWidget(time_box)

        return header

    # ========================================================
    # HOME
    # ========================================================

    def action_button(
        self,
        title,
        subtitle,
        kind
    ):

        btn = QPushButton()

        btn.setProperty(
            "kind",
            kind
        )

        btn.setCursor(
            Qt.PointingHandCursor
        )

        btn.setMinimumHeight(
            70
        )

        lay = QVBoxLayout(btn)

        lay.setContentsMargins(
            14,
            8,
            14,
            8
        )

        lay.setSpacing(0)

        ttl = QLabel(title)

        ttl.setObjectName(
            "actionTitle"
        )

        ttl.setAlignment(
            Qt.AlignCenter
        )

        ttl.setAttribute(
            Qt.WA_TransparentForMouseEvents
        )

        sub = QLabel(subtitle)

        sub.setObjectName(
            "actionSubtitle"
        )

        sub.setAlignment(
            Qt.AlignCenter
        )

        sub.setAttribute(
            Qt.WA_TransparentForMouseEvents
        )

        lay.addWidget(ttl)

        lay.addWidget(sub)

        return btn

    def build_home(self):

        page = QWidget()

        page.setObjectName(
            "page"
        )

        lay = QVBoxLayout(page)

        lay.setContentsMargins(
            16,
            14,
            16,
            14
        )

        lay.setSpacing(10)

        top = QHBoxLayout()

        top.setSpacing(12)

        counter = QFrame()

        counter.setObjectName(
            "card"
        )

        cl = QVBoxLayout(counter)

        cl.setContentsMargins(
            18,
            14,
            18,
            14
        )

        t = QLabel(
            "Current Carton Count"
        )

        t.setObjectName(
            "homeTitle"
        )

        t.setAlignment(
            Qt.AlignCenter
        )

        self.big_count = QLabel(
            "0"
        )

        self.big_count.setObjectName(
            "bigCount"
        )

        self.big_count.setAlignment(
            Qt.AlignCenter
        )

        self.goal_line = QLabel()

        self.goal_line.setObjectName(
            "goalLine"
        )

        self.goal_line.setAlignment(
            Qt.AlignCenter
        )

        cl.addWidget(t)

        cl.addWidget(
            self.big_count,
            1
        )

        cl.addWidget(
            self.goal_line
        )

        top.addWidget(
            counter,
            7
        )

        right = QVBoxLayout()

        right.setSpacing(10)

        grid = QGridLayout()

        grid.setSpacing(10)

        self.goal_card = InfoCard(
            "Target",
            "4,000",
            "cartons"
        )

        self.remaining_card = InfoCard(
            "Remaining",
            "4,000",
            "cartons"
        )

        progress_card = QFrame()

        progress_card.setObjectName(
            "card"
        )

        pl = QVBoxLayout(progress_card)

        pl.setContentsMargins(
            12,
            10,
            12,
            10
        )

        pt = QLabel(
            "Progress"
        )

        pt.setObjectName(
            "infoTitle"
        )

        pt.setAlignment(
            Qt.AlignCenter
        )

        self.donut = DonutWidget()

        pl.addWidget(pt)

        pl.addWidget(
            self.donut,
            1
        )

        self.status_card = InfoCard(
            "System Status",
            "Sensor not connected",
            "Session stopped"
        )

        grid.addWidget(
            self.goal_card,
            0,
            0
        )

        grid.addWidget(
            self.remaining_card,
            0,
            1
        )

        grid.addWidget(
            progress_card,
            1,
            0
        )

        grid.addWidget(
            self.status_card,
            1,
            1
        )

        right.addLayout(grid)

        actions = QGridLayout()

        actions.setSpacing(10)

        self.start_btn = self.action_button(
            "Start Session",
            "Begin recording",
            "green"
        )

        self.end_btn = self.action_button(
            "End Session",
            "Stop & save",
            "red"
        )

        actions.addWidget(
            self.start_btn,
            0,
            0
        )

        actions.addWidget(
            self.end_btn,
            0,
            1
        )

        right.addLayout(actions)

        # Quick +/- buttons
        steps = QHBoxLayout()

        steps.setSpacing(6)

        self.step_buttons = []

        for delta in [-10, -5, -1, 1, 5, 10]:

            b = QPushButton(
                f"{delta:+d}".replace("-", "−")
            )

            b.setObjectName(
                "stepPlus"
                if delta > 0
                else "stepMinus"
            )

            b.setCursor(
                Qt.PointingHandCursor
            )

            b.setMinimumHeight(
                48
            )

            b.clicked.connect(
                lambda _=False, d=delta:
                self.change_count(
                    d,
                    "Add" if d > 0 else "Remove"
                )
            )

            steps.addWidget(b)

            self.step_buttons.append(b)

        right.addLayout(steps)

        # Manual entry + resets
        tools = QHBoxLayout()

        tools.setSpacing(6)

        self.manual_btn = QPushButton(
            "Manual Entry"
        )

        self.reset_session_btn = QPushButton(
            "Reset Session"
        )

        self.reset_all_btn = QPushButton(
            "Reset All"
        )

        self.manual_btn.setObjectName(
            "toolButton"
        )

        self.reset_session_btn.setObjectName(
            "toolButton"
        )

        self.reset_all_btn.setObjectName(
            "dangerButton"
        )

        for b in [
            self.manual_btn,
            self.reset_session_btn,
            self.reset_all_btn
        ]:

            b.setCursor(
                Qt.PointingHandCursor
            )

            b.setMinimumHeight(
                44
            )

            tools.addWidget(b)

        right.addLayout(tools)

        top.addLayout(
            right,
            4
        )

        lay.addLayout(
            top,
            1
        )

        bottom = QHBoxLayout()

        bottom.setSpacing(10)

        self.duration_card = InfoCard(
            "Session Duration",
            "00:00:00",
            "Session inactive"
        )

        self.period_count_card = InfoCard(
            "Cartons This Session",
            "0",
            "cartons"
        )

        self.avg_time_card = InfoCard(
            "Time / Carton (EMA)",
            "0.00",
            f"seconds • N={EMA_N}"
        )

        self.countdown_card = InfoCard(
            "Work Period Remaining",
            "00:00:00",
            "Not started"
        )

        last = QFrame()

        last.setObjectName(
            "card"
        )

        ll = QVBoxLayout(last)

        ll.setContentsMargins(
            16,
            12,
            16,
            12
        )

        lt = QLabel(
            "Last Session Summary"
        )

        lt.setObjectName(
            "infoTitle"
        )

        lt.setAlignment(
            Qt.AlignCenter
        )

        self.last_summary = QLabel(
            "No previous session"
        )

        self.last_summary.setObjectName(
            "summary"
        )

        self.last_summary.setWordWrap(
            True
        )

        self.last_summary.setAlignment(
            Qt.AlignCenter
        )

        ll.addWidget(lt)

        ll.addWidget(
            self.last_summary,
            1
        )

        bottom.addWidget(
            self.duration_card
        )

        bottom.addWidget(
            self.countdown_card
        )

        bottom.addWidget(
            self.period_count_card
        )

        bottom.addWidget(
            self.avg_time_card
        )

        bottom.addWidget(last)

        lay.addLayout(bottom)

        return page

    # ========================================================
    # PAGE TITLE
    # ========================================================

    def page_title(
        self,
        text,
        subtitle
    ):

        w = QWidget()

        l = QVBoxLayout(w)

        l.setContentsMargins(
            0,
            0,
            0,
            0
        )

        l.setSpacing(2)

        t = QLabel(text)

        t.setObjectName(
            "pageTitle"
        )

        t.setAlignment(
            Qt.AlignCenter
        )

        s = QLabel(subtitle)

        s.setObjectName(
            "pageSubtitle"
        )

        s.setAlignment(
            Qt.AlignCenter
        )

        l.addWidget(t)

        l.addWidget(s)

        return w

    # ========================================================
    # REPORTS
    # ========================================================

    def build_reports(self):

        page = QWidget()

        page.setObjectName(
            "page"
        )

        lay = QVBoxLayout(page)

        lay.setContentsMargins(
            20,
            14,
            20,
            16
        )

        lay.setSpacing(10)

        lay.addWidget(
            self.page_title(
                "Reports",
                "Session summaries and performance"
            )
        )

        export_row = QHBoxLayout()

        export_row.addStretch(1)

        self.export_btn = QPushButton(
            "Export Excel"
        )

        self.export_btn.setObjectName(
            "primaryButton"
        )

        self.export_btn.setCursor(
            Qt.PointingHandCursor
        )

        self.export_btn.clicked.connect(
            self.export_excel
        )

        export_row.addWidget(
            self.export_btn
        )

        lay.addLayout(export_row)

        summary = QHBoxLayout()

        summary.setSpacing(10)

        self.report_total = InfoCard(
            "Total Cartons",
            "0",
            "cartons"
        )

        self.report_sessions = InfoCard(
            "Sessions",
            "0",
            "sessions"
        )

        self.report_avg = InfoCard(
            "Avg Time / Carton",
            "0.00",
            "seconds"
        )

        self.report_best = InfoCard(
            "Best Session",
            "0",
            "cartons"
        )

        for c in [
            self.report_total,
            self.report_sessions,
            self.report_avg,
            self.report_best
        ]:

            summary.addWidget(c)

        lay.addLayout(summary)

        body = QHBoxLayout()

        body.setSpacing(10)

        table_card = QFrame()

        table_card.setObjectName(
            "card"
        )

        tl = QVBoxLayout(table_card)

        tl.setContentsMargins(
            12,
            10,
            12,
            10
        )

        tt = QLabel(
            "Session Details"
        )

        tt.setObjectName(
            "panelTitle"
        )

        tt.setAlignment(
            Qt.AlignLeft
        )

        tl.addWidget(tt)

        self.reports_table = QTableWidget(
            0,
            7
        )

        self.reports_table.setHorizontalHeaderLabels(
            [
                "Session",
                "Start",
                "End",
                "Duration",
                "Cartons",
                "Goal",
                "Sec/Carton"
            ]
        )

        self.reports_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch
        )

        self.reports_table.setAlternatingRowColors(
            True
        )

        self.reports_table.setEditTriggers(
            QTableWidget.NoEditTriggers
        )

        tl.addWidget(
            self.reports_table,
            1
        )

        body.addWidget(
            table_card,
            7
        )

        chart_card = QFrame()

        chart_card.setObjectName(
            "card"
        )

        cl = QVBoxLayout(chart_card)

        cl.setContentsMargins(
            12,
            10,
            12,
            10
        )

        ct = QLabel(
            "Production by Day"
        )

        ct.setObjectName(
            "panelTitle"
        )

        ct.setAlignment(
            Qt.AlignLeft
        )

        self.bar_chart = BarChartWidget()

        cl.addWidget(ct)

        cl.addWidget(
            self.bar_chart,
            1
        )

        body.addWidget(
            chart_card,
            4
        )

        lay.addLayout(
            body,
            1
        )

        return page

    # ========================================================
    # LOG
    # ========================================================

    def build_log(self):

        page = QWidget()

        page.setObjectName(
            "page"
        )

        lay = QVBoxLayout(page)

        lay.setContentsMargins(
            20,
            14,
            20,
            16
        )

        lay.setSpacing(10)

        lay.addWidget(
            self.page_title(
                "Operation Log",
                "Additions, removals and session events"
            )
        )

        summary = QHBoxLayout()

        summary.setSpacing(10)

        self.log_count_card = InfoCard(
            "Operations Today",
            "0",
            "operations"
        )

        self.log_sync_card = InfoCard(
            "Last Sync",
            "Now",
            "data up to date"
        )

        summary.addWidget(
            self.log_count_card
        )

        summary.addWidget(
            self.log_sync_card
        )

        lay.addLayout(summary)

        row = QHBoxLayout()

        self.filter_combo = QComboBox()

        self.filter_combo.addItems(
            [
                "All",
                "Add",
                "Remove",
                "Start",
                "End",
                "Sync",
                "Settings"
            ]
        )

        self.filter_combo.currentTextChanged.connect(
            self.refresh_log
        )

        row.addStretch(1)

        row.addWidget(
            QLabel("Filter:")
        )

        row.addWidget(
            self.filter_combo
        )

        lay.addLayout(row)

        self.log_table = QTableWidget(
            0,
            3
        )

        self.log_table.setHorizontalHeaderLabels(
            [
                "Time",
                "Operation",
                "Details"
            ]
        )

        self.log_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch
        )

        self.log_table.setAlternatingRowColors(
            True
        )

        self.log_table.setEditTriggers(
            QTableWidget.NoEditTriggers
        )

        lay.addWidget(
            self.log_table,
            1
        )

        return page

    # ========================================================
    # SETTINGS
    # ========================================================

    def build_settings(self):

        page = QWidget()

        page.setObjectName(
            "page"
        )

        lay = QVBoxLayout(page)

        lay.setContentsMargins(
            20,
            14,
            20,
            16
        )

        lay.setSpacing(10)

        lay.addWidget(
            self.page_title(
                "Settings",
                "Configure target, display and sensor"
            )
        )

        grid = QGridLayout()

        grid.setSpacing(10)

        # ----------------------------------------------------
        # TARGET
        # ----------------------------------------------------

        target_card = QFrame()

        target_card.setObjectName(
            "card"
        )

        tc = QVBoxLayout(target_card)

        tc.setContentsMargins(
            20,
            18,
            20,
            18
        )

        tt = QLabel(
            "Target"
        )

        tt.setObjectName(
            "settingsTitle"
        )

        td = QLabel(
            "Total number of cartons to reach"
        )

        td.setObjectName(
            "settingsDesc"
        )

        self.target_spin = QSpinBox()

        self.target_spin.setRange(
            1,
            1000000
        )

        self.target_spin.setValue(
            self.state["target"]
        )

        self.target_spin.setMinimumHeight(
            48
        )

        tc.addWidget(tt)

        tc.addWidget(td)

        tc.addSpacing(10)

        tc.addWidget(
            self.target_spin
        )

        # ----------------------------------------------------
        # WORK PERIOD
        # ----------------------------------------------------

        period_card = QFrame()

        period_card.setObjectName(
            "card"
        )

        pc = QVBoxLayout(period_card)

        pc.setContentsMargins(
            20,
            18,
            20,
            18
        )

        pt = QLabel(
            "فترة الشغل  (Work Period)"
        )

        pt.setObjectName(
            "settingsTitle"
        )

        pd = QLabel(
            "مدة فترة الشغل بالدقايق — "
            "الكاونتداون يبدأ مع Start Session"
        )

        pd.setObjectName(
            "settingsDesc"
        )

        pd.setWordWrap(True)

        self.period_spin = QSpinBox()

        self.period_spin.setRange(
            1,
            1440
        )

        self.period_spin.setSuffix(
            " min"
        )

        self.period_spin.setValue(
            self.state.get(
                "work_period_minutes",
                DEFAULT_WORK_MINUTES
            )
        )

        self.period_spin.setMinimumHeight(
            48
        )

        pc.addWidget(pt)

        pc.addWidget(pd)

        pc.addSpacing(10)

        pc.addWidget(
            self.period_spin
        )

        # ----------------------------------------------------
        # DISPLAY
        # ----------------------------------------------------

        display_card = QFrame()

        display_card.setObjectName(
            "card"
        )

        dc = QVBoxLayout(display_card)

        dc.setContentsMargins(
            20,
            18,
            20,
            18
        )

        dt = QLabel(
            "Display"
        )

        dt.setObjectName(
            "settingsTitle"
        )

        dd = QLabel(
            "Fullscreen options"
        )

        dd.setObjectName(
            "settingsDesc"
        )

        self.fullscreen_check = QCheckBox(
            "Dashboard fullscreen"
        )

        self.counter_fullscreen_check = QCheckBox(
            "Counter screen fullscreen"
        )

        self.counter_fullscreen_check.setChecked(
            True
        )

        self.show_pie_check = QCheckBox(
            "Show pie chart on display"
        )

        self.show_pie_check.setChecked(
            bool(self.state.get("show_pie", True))
        )

        # Applies right away, no need to press Save
        self.show_pie_check.toggled.connect(
            self.set_show_pie
        )

        dc.addWidget(dt)

        dc.addWidget(dd)

        dc.addSpacing(10)

        dc.addWidget(
            self.fullscreen_check
        )

        dc.addWidget(
            self.counter_fullscreen_check
        )

        dc.addWidget(
            self.show_pie_check
        )

        dc.addStretch(1)

        # ----------------------------------------------------
        # SENSOR
        # ----------------------------------------------------

        sensor_card = QFrame()

        sensor_card.setObjectName(
            "card"
        )

        sc = QVBoxLayout(sensor_card)

        sc.setContentsMargins(
            20,
            18,
            20,
            18
        )

        st = QLabel(
            "Sensor — TTL to USB"
        )

        st.setObjectName(
            "settingsTitle"
        )

        sd = QLabel(
            "Connect the TTL-to-USB converter to the PC, "
            "select its COM port, baud rate and counting mode."
        )

        sd.setObjectName(
            "settingsDesc"
        )

        sd.setWordWrap(True)

        cfg = dict(
            DEFAULT_SERIAL,
            **self.state.get(
                "serial",
                {}
            )
        )

        form = QGridLayout()

        form.setHorizontalSpacing(10)

        form.setVerticalSpacing(6)

        # COM port
        self.port_combo = QComboBox()

        refresh_ports = QPushButton(
            "Refresh"
        )

        refresh_ports.setObjectName(
            "smallButton"
        )

        refresh_ports.setToolTip(
            "Refresh COM ports"
        )

        refresh_ports.clicked.connect(
            self.reload_ports
        )

        port_row = QHBoxLayout()

        port_row.addWidget(
            self.port_combo,
            1
        )

        port_row.addWidget(
            refresh_ports
        )

        form.addWidget(
            QLabel("COM Port"),
            0,
            0
        )

        form.addLayout(
            port_row,
            0,
            1
        )

        # Baud
        self.baud_combo = QComboBox()

        for b in BAUD_CHOICES:

            self.baud_combo.addItem(
                str(b),
                b
            )

        form.addWidget(
            QLabel("Baud"),
            1,
            0
        )

        form.addWidget(
            self.baud_combo,
            1,
            1
        )

        # Mode
        self.mode_combo = QComboBox()

        for key, label in INPUT_MODES:

            self.mode_combo.addItem(
                label,
                key
            )

        form.addWidget(
            QLabel("Counting Mode"),
            2,
            0
        )

        form.addWidget(
            self.mode_combo,
            2,
            1
        )

        # Debounce
        self.debounce_spin = QSpinBox()

        self.debounce_spin.setRange(
            0,
            2000
        )

        self.debounce_spin.setSuffix(
            " ms"
        )

        form.addWidget(
            QLabel("Debounce"),
            3,
            0
        )

        form.addWidget(
            self.debounce_spin,
            3,
            1
        )

        form.setColumnStretch(
            1,
            1
        )

        self.sensor_state_label = QLabel(
            "Sensor not connected"
        )

        self.sensor_state_label.setObjectName(
            "sensorState"
        )

        self.sensor_state_label.setWordWrap(
            True
        )

        sc.addWidget(st)

        sc.addWidget(sd)

        sc.addSpacing(8)

        sc.addLayout(form)

        sc.addSpacing(8)

        sc.addWidget(
            self.sensor_state_label
        )

        sc.addStretch(1)

        self.load_serial_controls(
            cfg
        )

        grid.addWidget(
            target_card,
            0,
            0
        )

        grid.addWidget(
            period_card,
            0,
            1
        )

        grid.addWidget(
            display_card,
            1,
            0
        )

        grid.addWidget(
            sensor_card,
            1,
            1
        )

        lay.addLayout(
            grid,
            1
        )

        buttons = QHBoxLayout()

        wipe = QPushButton(
            "Reset All App Data"
        )

        wipe.setObjectName(
            "dangerButton"
        )

        wipe.setMinimumHeight(
            44
        )

        wipe.setCursor(
            Qt.PointingHandCursor
        )

        wipe.clicked.connect(
            lambda: self.reset_all(wipe_settings=True)
        )

        buttons.addWidget(wipe)

        buttons.addStretch(1)

        reset = QPushButton(
            "Reset Defaults"
        )

        reset.setObjectName(
            "secondaryButton"
        )

        reset.clicked.connect(
            self.reset_settings
        )

        save = QPushButton(
            "Save Settings"
        )

        save.setObjectName(
            "primaryButton"
        )

        save.clicked.connect(
            self.save_settings
        )

        buttons.addWidget(reset)

        buttons.addWidget(save)

        lay.addLayout(buttons)

        return page

    # ========================================================
    # ACTIONS
    # ========================================================

    def connect_actions(self):

        self.nav_home.clicked.connect(
            lambda:
            self.switch_page(
                0,
                self.nav_home
            )
        )

        self.nav_reports.clicked.connect(
            lambda:
            self.switch_page(
                1,
                self.nav_reports
            )
        )

        self.nav_log.clicked.connect(
            lambda:
            self.switch_page(
                2,
                self.nav_log
            )
        )

        self.nav_settings.clicked.connect(
            lambda:
            self.switch_page(
                3,
                self.nav_settings
            )
        )

        self.nav_exit.clicked.connect(
            self.close
        )

        self.start_btn.clicked.connect(
            self.start_session
        )

        self.end_btn.clicked.connect(
            self.end_session
        )

        self.manual_btn.clicked.connect(
            self.manual_entry
        )

        self.reset_session_btn.clicked.connect(
            self.reset_session
        )

        self.reset_all_btn.clicked.connect(
            lambda: self.reset_all(wipe_settings=False)
        )

    # ========================================================
    # PAGE SWITCH
    # ========================================================

    def switch_page(
        self,
        index,
        button
    ):

        self.pages.setCurrentIndex(
            index
        )

        for b in [
            self.nav_home,
            self.nav_reports,
            self.nav_log,
            self.nav_settings
        ]:

            b.setChecked(
                b is button
            )

        if index == 1:

            self.refresh_reports()

        elif index == 2:

            self.refresh_log()

    # ========================================================
    # SENSOR STATUS
    # ========================================================

    def on_sensor_status(
        self,
        text,
        connected
    ):

        self.sensor_connected = connected

        self.usb_label.setText(
            text
        )

        self.usb_label.setProperty(
            "connected",
            connected
        )

        self.usb_label.style().unpolish(
            self.usb_label
        )

        self.usb_label.style().polish(
            self.usb_label
        )

        self.sensor_state_label.setText(
            text
        )

        self.sensor_state_label.setStyleSheet(
            f"""
            color: {GREEN if connected else RED};
            font-family: {FONT};
            font-size: 18px;
            font-weight: 900;
            """
        )

        self.refresh_all()

    # ========================================================
    # SENSOR DELTA
    # ========================================================

    def on_sensor_delta(
        self,
        delta
    ):

        self.change_count(
            delta,
            "Add"
        )

    # ========================================================
    # CHANGE COUNT
    # ========================================================

    def change_count(
        self,
        delta,
        operation
    ):

        old = self.state["count"]

        self.state["count"] = max(
            0,
            self.state["count"] + int(delta)
        )

        actual = (
            self.state["count"]
            - old
        )

        if actual != 0:

            self.log_operation(
                operation,
                f"{actual:+d} cartons | "
                f"total {self.state['count']:,}"
            )

        # Only single cartons are real cycle-time samples;
        # bulk +5/+10/manual edits are not timed.
        if self.session_active and actual == 1:

            self._record_carton()

        self.save_state()

        self.refresh_all()

        self.count_changed.emit(
            self.state["count"]
        )

    # ========================================================
    # EMA OF TIME PER CARTON
    # ========================================================

    def ideal_seconds(self):
        # Pace needed to hit the session goal in the session time
        if not self.session_active or self.session_goal <= 0:
            return 0.0

        return (
            self.session_minutes * 60
            / self.session_goal
        )

    def _reset_ema(self):

        self.ema_value = None

        self.ema_samples = 0

        self.last_carton_mono = None

        self.session_carton_index = 0

    def _record_carton(self):

        now = time.monotonic()

        self.session_carton_index += 1

        cycle = None

        if self.last_carton_mono is not None:

            cycle = now - self.last_carton_mono

            self.ema_samples += 1

            if self.ema_samples <= EMA_N:

                # Bootstrap: simple average of the first N cycles
                prev = self.ema_value or 0.0

                self.ema_value = (
                    prev * (self.ema_samples - 1) + cycle
                ) / self.ema_samples

            else:

                alpha = 2 / (EMA_N + 1)

                self.ema_value = (
                    alpha * cycle
                    + (1 - alpha) * self.ema_value
                )

        self.last_carton_mono = now

        DATA_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        new_file = not CARTONS_FILE.exists()

        try:

            with CARTONS_FILE.open(
                "a",
                newline="",
                encoding="utf-8-sig" if new_file else "utf-8"
            ) as f:

                w = csv.writer(f)

                if new_file:

                    w.writerow(
                        [
                            "session_id",
                            "session_name",
                            "carton",
                            "time",
                            "cycle_seconds",
                            "ema_seconds",
                        ]
                    )

                w.writerow(
                    [
                        self.session_start_wall.isoformat(
                            timespec="microseconds"
                        ),
                        self.session_name,
                        self.session_carton_index,
                        datetime.now().isoformat(
                            timespec="milliseconds"
                        ),
                        "" if cycle is None else f"{cycle:.3f}",
                        ""
                        if self.ema_value is None
                        else f"{self.ema_value:.4f}",
                    ]
                )

        except Exception:
            pass

    # ========================================================
    # START
    # ========================================================

    def start_session(self):

        if self.session_active:
            return

        defaults = self.state.get(
            "session_defaults"
        ) or {}

        dlg = SessionDialog(
            self,
            name=defaults.get("name", ""),
            goal=int(defaults.get("goal", 500)),
            minutes=int(
                defaults.get(
                    "minutes",
                    self.state.get(
                        "work_period_minutes",
                        DEFAULT_WORK_MINUTES
                    )
                )
            )
        )

        if dlg.exec() != QDialog.Accepted:
            return

        name, goal, minutes = dlg.values()

        self.state["session_defaults"] = {
            "name": name,
            "goal": goal,
            "minutes": minutes,
        }

        self.session_name = name

        self.session_goal = goal

        self.session_minutes = minutes

        self._reset_ema()

        self.session_active = True

        self.session_start_count = (
            self.state["count"]
        )

        self.session_start_wall = (
            datetime.now()
        )

        self.session_start_mono = (
            time.monotonic()
        )

        self._sync_active_session()

        self.log_operation(
            "Start",
            f"Session '{name or '-'}' started at total "
            f"{self.state['count']:,} | "
            f"goal {goal:,} | {minutes} min"
        )

        self.save_state()

        self.refresh_all()

        self.count_changed.emit(
            self.state["count"]
        )

    def _sync_active_session(self):
        # Mirror the running session into state so the
        # display window can draw its goal pie chart.
        if self.session_active:

            self.state["active_session"] = {
                "name": self.session_name,
                "goal": self.session_goal,
                "minutes": self.session_minutes,
                "start_count": self.session_start_count,
            }

        else:

            self.state["active_session"] = None

    # ========================================================
    # END
    # ========================================================

    def end_session(self):

        if not self.session_active:
            return

        elapsed = max(
            0.0,
            time.monotonic()
            - self.session_start_mono
        )

        cartons = max(
            0,
            self.state["count"]
            - self.session_start_count
        )

        sec_per = (
            elapsed / cartons
            if cartons > 0
            else 0.0
        )

        end = datetime.now()

        self.state["last_session"] = {
            "start":
                self.session_start_wall.isoformat(
                    timespec="seconds"
                ),

            "end":
                end.isoformat(
                    timespec="seconds"
                ),

            "duration_seconds":
                elapsed,

            "cartons":
                cartons,

            "seconds_per_carton":
                sec_per,

            "name":
                self.session_name,

            "goal":
                self.session_goal,

            "planned_minutes":
                self.session_minutes,

            "ema_seconds_per_carton":
                self.ema_value or 0.0,

            "full_target":
                self.state["target"],

            # Unique key linking to carton_events.csv
            "session_id":
                self.session_start_wall.isoformat(
                    timespec="microseconds"
                ),
        }

        self.write_session(
            self.state["last_session"]
        )

        self.session_active = False

        self.session_start_mono = None

        self.session_start_wall = None

        self._sync_active_session()

        goal = self.session_goal

        reached = cartons >= goal

        self.log_operation(
            "End",
            f"{cartons:,}/{goal:,} cartons | "
            f"{fmt_hms(elapsed)} | "
            f"{sec_per:.2f} sec/carton"
        )

        self.save_state()

        self.refresh_all()

        self.count_changed.emit(
            self.state["count"]
        )

        QMessageBox.information(
            self,
            "Session Summary",
            f"Session: {self.session_name or '-'}\n"
            f"Duration: {fmt_hms(elapsed)}\n"
            f"Cartons: {cartons:,} / goal {goal:,}"
            f" ({round(cartons / max(1, goal) * 100)}%)\n"
            f"{'Goal reached ✔' if reached else 'Goal not reached'}\n"
            f"Avg time/carton: {sec_per:.2f} seconds\n"
            f"EMA time/carton (N={EMA_N}): "
            f"{self.ema_value or 0.0:.2f} seconds\n"
            f"Ideal time/carton: "
            f"{self.session_minutes * 60 / max(1, goal):.2f} seconds"
        )

    # ========================================================
    # RESET SESSION
    # ========================================================

    def reset_session(self):

        if not self.session_active:

            QMessageBox.information(
                self,
                "Reset Session",
                "No session is running."
            )

            return

        res = QMessageBox.question(
            self,
            "Reset Session",
            "Restart the current session?\n\n"
            "Session cartons go back to 0 and the timer "
            "restarts.\nThe overall total is kept.",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if res != QMessageBox.Yes:
            return

        self.session_start_count = (
            self.state["count"]
        )

        self.session_start_wall = (
            datetime.now()
        )

        self.session_start_mono = (
            time.monotonic()
        )

        self._reset_ema()

        self._sync_active_session()

        self.log_operation(
            "Reset Session",
            f"Session restarted at total "
            f"{self.state['count']:,}"
        )

        self.save_state()

        self.refresh_all()

        self.count_changed.emit(
            self.state["count"]
        )

    # ========================================================
    # RESET ALL
    # ========================================================

    def reset_all(self, wipe_settings=False):

        res = QMessageBox.warning(
            self,
            "Reset All App Data" if wipe_settings else "Reset All",
            "Reset the whole app?\n\n"
            "• Total count goes back to 0\n"
            "• Running session is cancelled (not saved)\n"
            "• Session history, carton times and operation "
            "log are cleared\n"
            + (
                "• Settings go back to defaults "
                "(target, work period, last session values)\n\n"
                "A backup is saved in the data/backup folder."
                if wipe_settings
                else
                "\nSettings are kept, and a backup of the data "
                "is saved in the data/backup folder."
            ),
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if res != QMessageBox.Yes:
            return

        # Back up before wiping anything
        backup = (
            DATA_DIR
            / "backup"
            / datetime.now().strftime("%Y%m%d_%H%M%S")
        )

        # Never overwrite an earlier backup from the same second
        n = 2

        while backup.exists():

            backup = backup.with_name(
                f"{backup.name.split('-')[0]}-{n}"
            )

            n += 1

        try:

            backup.mkdir(
                parents=True,
                exist_ok=True
            )

            for f in [
                STATE_FILE,
                SESSIONS_FILE,
                LOG_FILE,
                CARTONS_FILE
            ]:

                if f.exists():

                    (backup / f.name).write_bytes(
                        f.read_bytes()
                    )

        except Exception as e:

            QMessageBox.critical(
                self,
                "Reset All",
                f"Backup failed, nothing was reset.\n\n{e}"
            )

            return

        self.session_active = False

        self.session_start_mono = None

        self.session_start_wall = None

        self._reset_ema()

        self._sync_active_session()

        self.state["count"] = 0

        self.state["last_session"] = None

        if wipe_settings:

            self.state["target"] = DEFAULT_TARGET

            self.state["work_period_minutes"] = DEFAULT_WORK_MINUTES

            self.state["serial"] = dict(DEFAULT_SERIAL)

            self.state.pop("session_defaults", None)

            # Put the Settings page widgets back to defaults too
            self.reset_settings()

        for f in [
            SESSIONS_FILE,
            LOG_FILE,
            CARTONS_FILE
        ]:

            try:
                f.unlink(missing_ok=True)
            except Exception:
                pass

        self.log_operation(
            "Reset All",
            f"{'All app data and settings' if wipe_settings else 'App data'} "
            f"reset (backup: {backup.name})"
        )

        self.save_state()

        self.refresh_all()

        self.refresh_reports()

        self.refresh_log()

        self.count_changed.emit(
            self.state["count"]
        )

        QMessageBox.information(
            self,
            "Reset All",
            f"Done. Backup saved to:\n{backup}"
        )

    # ========================================================
    # EXPORT EXCEL
    # ========================================================

    def export_excel(self):

        try:
            import openpyxl  # noqa: F401
        except ImportError:

            QMessageBox.warning(
                self,
                "Export Excel",
                "openpyxl is not installed.\n\n"
                "Run: pip install openpyxl"
            )

            return

        default = (
            Path.home()
            / "Documents"
            / f"WorkingCamp_{datetime.now():%Y%m%d_%H%M}.xlsx"
        )

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Excel",
            str(default),
            "Excel (*.xlsx)"
        )

        if not path:
            return

        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        try:

            export_workbook(
                path,
                self.read_sessions(),
                read_carton_events(),
                self.state["target"]
            )

        except PermissionError:

            QMessageBox.warning(
                self,
                "Export Excel",
                "Couldn't write the file.\n"
                "Close it in Excel and try again."
            )

            return

        except Exception as e:

            QMessageBox.critical(
                self,
                "Export Excel",
                f"Export failed:\n{e}"
            )

            return

        self.log_operation(
            "Export",
            f"Excel report saved: {path}"
        )

        QMessageBox.information(
            self,
            "Export Excel",
            f"Saved:\n{path}"
        )

    # ========================================================
    # MANUAL ENTRY
    # ========================================================

    def manual_entry(self):

        dlg = ManualEntryDialog(
            self,
            self.state["count"]
        )

        if dlg.exec() != QDialog.Accepted:
            return

        mode, amount = dlg.values()

        if mode == "add":

            delta = amount

        elif mode == "subtract":

            delta = -amount

        else:

            delta = amount - self.state["count"]

        if delta == 0:
            return

        self.change_count(
            delta,
            "Manual"
        )

    # ========================================================
    # SESSION STATS
    # ========================================================

    def current_session_stats(self):

        if self.session_active:

            elapsed = max(
                0.0,
                time.monotonic()
                - self.session_start_mono
            )

            cartons = max(
                0,
                self.state["count"]
                - self.session_start_count
            )

            sec_per = (
                elapsed / cartons
                if cartons > 0
                else 0.0
            )

            return (
                elapsed,
                cartons,
                sec_per,
                "Session active"
            )

        return (
            0.0,
            0,
            0.0,
            "Session inactive"
        )

    # ========================================================
    # COUNTDOWN
    # ========================================================

    def get_countdown(self):

        if not self.session_active:

            return (
                0.0,
                False,
                False
            )

        elapsed = max(
            0.0,
            time.monotonic()
            - self.session_start_mono
        )

        total = max(
            1,
            self.session_minutes
        ) * 60

        remaining = (
            total
            - elapsed
        )

        return (
            max(
                0.0,
                remaining
            ),
            remaining <= 0,
            True
        )

    # ========================================================
    # REFRESH
    # ========================================================

    def refresh_all(self):

        count = self.state["count"]

        target = self.state["target"]

        self.big_count.setText(
            f"{count:,}"
        )

        self.goal_line.setText(
            f"of {target:,} cartons"
            f" • "
            f"{round(min(1, count / target) * 100)}%"
        )

        self.donut.set_values(
            count,
            target
        )

        self.goal_card.value.setText(
            f"{target:,}"
        )

        self.remaining_card.value.setText(
            f"{max(0, target - count):,}"
        )

        self.status_card.value.setText(
            "Sensor connected"
            if self.sensor_connected
            else "Sensor not connected"
        )

        self.status_card.value.setStyleSheet(
            f"color: "
            f"{GREEN if self.sensor_connected else RED};"
        )

        self.status_card.subtitle.setText(
            "Session active"
            if self.session_active
            else "Session inactive"
        )

        elapsed, cartons, sec_per, text = (
            self.current_session_stats()
        )

        self.duration_card.value.setText(
            fmt_hms(elapsed)
        )

        self.duration_card.subtitle.setText(
            text
        )

        self.period_count_card.value.setText(
            f"{cartons:,}"
        )

        if self.session_active:

            self.period_count_card.subtitle.setText(
                f"of {self.session_goal:,} goal • "
                f"{round(cartons / max(1, self.session_goal) * 100)}%"
            )

        else:

            self.period_count_card.subtitle.setText(
                "cartons"
            )

        self.avg_time_card.value.setText(
            f"{self.ema_value or 0.0:.2f}"
        )

        self.avg_time_card.subtitle.setText(
            f"Avg {sec_per:.2f} • "
            f"Ideal {self.ideal_seconds():.2f} sec"
        )

        self.avg_time_card.value.setStyleSheet(
            f"color: {RED};"
            if self.session_active
            and self.ema_value
            and self.ema_value > self.ideal_seconds() > 0
            else ""
        )

        remaining, over, active = (
            self.get_countdown()
        )

        if not active:

            self.countdown_card.value.setText(
                fmt_hms(
                    self.state.get(
                        "work_period_minutes",
                        DEFAULT_WORK_MINUTES
                    ) * 60
                )
            )

            self.countdown_card.subtitle.setText(
                "Not started"
            )

            self.countdown_card.value.setStyleSheet(
                f"color: {NAVY};"
            )

        elif over:

            self.countdown_card.value.setText(
                "00:00:00"
            )

            self.countdown_card.subtitle.setText(
                "Work period ended"
            )

            self.countdown_card.value.setStyleSheet(
                f"color: {RED};"
            )

        else:

            self.countdown_card.value.setText(
                fmt_hms(remaining)
            )

            self.countdown_card.subtitle.setText(
                "Counting down"
            )

            self.countdown_card.value.setStyleSheet(
                f"color: {NAVY};"
            )

        last = self.state.get(
            "last_session"
        )

        if last:

            self.last_summary.setText(
                f"Duration: "
                f"{fmt_hms(last.get('duration_seconds', 0))}\n"
                f"Cartons: "
                f"{int(last.get('cartons', 0)):,}"
                + (
                    f" / {int(last['goal']):,}"
                    if last.get("goal")
                    else ""
                )
                + "\n"
                f"Time/carton: "
                f"{float(last.get('seconds_per_carton', 0)):.2f} sec"
            )

        else:

            self.last_summary.setText(
                "No previous session"
            )

        self.start_btn.setEnabled(
            not self.session_active
        )

        self.end_btn.setEnabled(
            self.session_active
        )

        self.reset_session_btn.setEnabled(
            self.session_active
        )

    # ========================================================
    # CLOCK
    # ========================================================

    def tick(self):

        now = datetime.now()

        self.date_label.setText(
            now.strftime("%d/%m/%Y")
        )

        self.time_label.setText(
            now.strftime("%H:%M")
        )

        remaining, over, active = (
            self.get_countdown()
        )

        self.countdown_changed.emit(
            int(remaining),
            over,
            active
        )

        _, cartons, sec_per, _ = (
            self.current_session_stats()
        )

        self.rate_changed.emit(
            float(self.ema_value or 0.0),
            float(sec_per),
            float(self.ideal_seconds()),
            int(self.ema_samples)
        )

        if self.session_active:

            self.refresh_all()

    # ========================================================
    # REPORTS
    # ========================================================

    def read_sessions(self):

        if not SESSIONS_FILE.exists():
            return []

        try:

            with SESSIONS_FILE.open(
                "r",
                encoding="utf-8-sig",
                newline=""
            ) as f:

                return list(
                    csv.DictReader(f)
                )

        except Exception:

            return []

    def refresh_reports(self):

        rows = self.read_sessions()

        self.reports_table.setRowCount(
            0
        )

        total = 0

        secs = []

        best = 0

        day_totals = defaultdict(int)

        for row in rows:

            cartons = int(
                float(
                    row.get(
                        "cartons",
                        0
                    ) or 0
                )
            )

            sec_per = float(
                row.get(
                    "seconds_per_carton",
                    0
                ) or 0
            )

            total += cartons

            best = max(
                best,
                cartons
            )

            if cartons > 0:

                secs.append(
                    sec_per
                )

            try:

                dt = datetime.fromisoformat(
                    row.get(
                        "start_time",
                        ""
                    )
                )

                day_totals[
                    dt.strftime("%d/%m")
                ] += cartons

            except Exception:
                pass

        for row in reversed(
            rows[-50:]
        ):

            i = self.reports_table.rowCount()

            self.reports_table.insertRow(
                i
            )

            try:

                start = datetime.fromisoformat(
                    row["start_time"]
                ).strftime(
                    "%d/%m/%Y %H:%M"
                )

            except Exception:

                start = row.get(
                    "start_time",
                    ""
                )

            try:

                end = datetime.fromisoformat(
                    row["end_time"]
                ).strftime(
                    "%d/%m/%Y %H:%M"
                )

            except Exception:

                end = row.get(
                    "end_time",
                    ""
                )

            duration = fmt_hms(
                float(
                    row.get(
                        "duration_seconds",
                        0
                    ) or 0
                )
            )

            cartons = int(
                float(
                    row.get(
                        "cartons",
                        0
                    ) or 0
                )
            )

            sec_per = float(
                row.get(
                    "seconds_per_carton",
                    0
                ) or 0
            )

            goal = row.get("goal") or ""

            if goal:

                try:
                    goal_i = int(float(goal))
                    goal = (
                        f"{goal_i:,} "
                        f"({round(cartons / max(1, goal_i) * 100)}%)"
                    )
                except ValueError:
                    pass

            for col, val in enumerate(
                [
                    row.get("name") or "-",
                    start,
                    end,
                    duration,
                    f"{cartons:,}",
                    goal or "-",
                    f"{sec_per:.2f}"
                ]
            ):

                item = QTableWidgetItem(
                    val
                )

                item.setTextAlignment(
                    Qt.AlignCenter
                )

                self.reports_table.setItem(
                    i,
                    col,
                    item
                )

        self.report_total.value.setText(
            f"{total:,}"
        )

        self.report_sessions.value.setText(
            str(len(rows))
        )

        self.report_avg.value.setText(
            f"{sum(secs) / len(secs):.2f}"
            if secs
            else "0.00"
        )

        self.report_best.value.setText(
            f"{best:,}"
        )

        self.bar_chart.set_data(
            list(day_totals.keys()),
            list(day_totals.values())
        )

    # ========================================================
    # LOG
    # ========================================================

    def log_operation(
        self,
        operation,
        details
    ):

        DATA_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        new_file = not LOG_FILE.exists()

        try:

            with LOG_FILE.open(
                "a",
                newline="",
                encoding="utf-8-sig"
            ) as f:

                w = csv.writer(f)

                if new_file:

                    w.writerow(
                        [
                            "time",
                            "operation",
                            "details"
                        ]
                    )

                w.writerow(
                    [
                        datetime.now().strftime(
                            "%Y-%m-%d %H:%M:%S"
                        ),
                        operation,
                        details
                    ]
                )

        except Exception:
            pass

        if hasattr(
            self,
            "log_table"
        ):

            self.refresh_log()

    def refresh_log(self):

        self.log_table.setRowCount(
            0
        )

        if not LOG_FILE.exists():

            self.log_count_card.value.setText(
                "0"
            )

            return

        try:

            with LOG_FILE.open(
                "r",
                encoding="utf-8-sig",
                newline=""
            ) as f:

                rows = list(
                    csv.DictReader(f)
                )

        except Exception:

            rows = []

        selected = (
            self.filter_combo.currentText()
        )

        def include(row):

            if selected == "All":
                return True

            return (
                selected.lower()
                in row.get(
                    "operation",
                    ""
                ).lower()
            )

        filtered = [
            r
            for r in rows
            if include(r)
        ]

        for row in reversed(
            filtered[-300:]
        ):

            i = self.log_table.rowCount()

            self.log_table.insertRow(
                i
            )

            for col, val in enumerate(
                [
                    row.get(
                        "time",
                        ""
                    ),
                    row.get(
                        "operation",
                        ""
                    ),
                    row.get(
                        "details",
                        ""
                    )
                ]
            ):

                item = QTableWidgetItem(
                    val
                )

                item.setTextAlignment(
                    Qt.AlignCenter
                )

                self.log_table.setItem(
                    i,
                    col,
                    item
                )

        today = datetime.now().strftime(
            "%Y-%m-%d"
        )

        today_count = sum(
            1
            for r in rows
            if r.get(
                "time",
                ""
            ).startswith(today)
        )

        self.log_count_card.value.setText(
            str(today_count)
        )

    # ========================================================
    # COM PORT SETTINGS
    # ========================================================

    def reload_ports(
        self,
        selected=None
    ):

        if selected is None:

            selected = (
                self.port_combo.currentData()
                or "auto"
            )

        self.port_combo.clear()

        self.port_combo.addItem(
            "Auto — Detect USB/TTL adapter",
            "auto"
        )

        ports = list_serial_ports()

        for dev, desc in ports:

            self.port_combo.addItem(
                f"{dev} — {desc}",
                dev
            )

        i = self.port_combo.findData(
            selected
        )

        if (
            i < 0
            and selected != "auto"
        ):

            self.port_combo.addItem(
                f"{selected} — Not currently plugged in",
                selected
            )

            i = (
                self.port_combo.count()
                - 1
            )

        self.port_combo.setCurrentIndex(
            max(
                0,
                i
            )
        )

    def load_serial_controls(
        self,
        cfg
    ):

        self.reload_ports(
            cfg.get(
                "port",
                "auto"
            )
        )

        baud = int(
            cfg.get(
                "baud",
                9600
            )
        )

        i = self.baud_combo.findData(
            baud
        )

        if i < 0:

            self.baud_combo.addItem(
                str(baud),
                baud
            )

            i = (
                self.baud_combo.count()
                - 1
            )

        self.baud_combo.setCurrentIndex(
            i
        )

        mode = cfg.get(
            "mode",
            "count_total"
        )

        mode_index = (
            self.mode_combo.findData(
                mode
            )
        )

        self.mode_combo.setCurrentIndex(
            max(
                0,
                mode_index
            )
        )

        self.debounce_spin.setValue(
            int(
                cfg.get(
                    "debounce_ms",
                    80
                )
            )
        )

    def current_serial_config(self):

        return {
            "port":
                self.port_combo.currentData()
                or "auto",

            "baud":
                int(
                    self.baud_combo.currentData()
                    or 9600
                ),

            "mode":
                self.mode_combo.currentData()
                or "count_total",

            "debounce_ms":
                int(
                    self.debounce_spin.value()
                ),
        }

    # ========================================================
    # SAVE SETTINGS
    # ========================================================

    def save_settings(self):

        self.state["target"] = (
            self.target_spin.value()
        )

        self.state["work_period_minutes"] = (
            self.period_spin.value()
        )

        old_serial = self.state.get(
            "serial"
        )

        new_serial = (
            self.current_serial_config()
        )

        self.state["serial"] = new_serial

        if new_serial != old_serial:

            self.serial_config_changed.emit(
                new_serial
            )

            self.log_operation(
                "Settings",
                f"Sensor: "
                f"{new_serial['port']} @ "
                f"{new_serial['baud']} • "
                f"{new_serial['mode']}"
            )

        if self.fullscreen_check.isChecked():

            self.showFullScreen()

        else:

            self.showNormal()

        self.log_operation(
            "Settings",
            f"Target changed to "
            f"{self.state['target']:,} | "
            f"Work period changed to "
            f"{self.state['work_period_minutes']} min"
        )

        self.save_state()

        self.refresh_all()

        self.count_changed.emit(
            self.state["count"]
        )

        QMessageBox.information(
            self,
            "Settings",
            "Settings saved."
        )

    # ========================================================
    # RESET
    # ========================================================

    def reset_settings(self):

        self.target_spin.setValue(
            DEFAULT_TARGET
        )

        self.period_spin.setValue(
            DEFAULT_WORK_MINUTES
        )

        self.load_serial_controls(
            DEFAULT_SERIAL
        )

        self.fullscreen_check.setChecked(
            False
        )

        self.counter_fullscreen_check.setChecked(
            True
        )

        self.show_pie_check.setChecked(
            True
        )

    def set_show_pie(self, show):

        self.state["show_pie"] = bool(show)

        self.save_state()

        self.count_changed.emit(
            self.state["count"]
        )

    # ========================================================
    # STATE
    # ========================================================

    def save_state(self):

        DATA_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        try:

            STATE_FILE.write_text(
                json.dumps(
                    self.state,
                    ensure_ascii=False,
                    indent=2
                ),
                encoding="utf-8"
            )

        except Exception:
            pass

    # ========================================================
    # SESSION FILE
    # ========================================================

    def write_session(
        self,
        session
    ):

        DATA_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        header = [
            "start_time",
            "end_time",
            "duration_seconds",
            "cartons",
            "seconds_per_carton",
            "name",
            "goal",
            "planned_minutes",
            "ema_seconds_per_carton",
            "full_target",
            "session_id",
        ]

        try:

            # Upgrade files written before name/goal columns existed
            if SESSIONS_FILE.exists():

                with SESSIONS_FILE.open(
                    "r",
                    newline="",
                    encoding="utf-8-sig"
                ) as f:

                    reader = csv.DictReader(f)

                    old_rows = list(reader)

                    old_header = reader.fieldnames or []

                if old_header != header:

                    with SESSIONS_FILE.open(
                        "w",
                        newline="",
                        encoding="utf-8-sig"
                    ) as f:

                        w = csv.DictWriter(
                            f,
                            fieldnames=header,
                            extrasaction="ignore"
                        )

                        w.writeheader()

                        w.writerows(old_rows)

            new_file = not SESSIONS_FILE.exists()

            with SESSIONS_FILE.open(
                "a",
                newline="",
                encoding="utf-8-sig" if new_file else "utf-8"
            ) as f:

                w = csv.writer(f)

                if new_file:

                    w.writerow(header)

                w.writerow(
                    [
                        session["start"],
                        session["end"],
                        f'{session["duration_seconds"]:.2f}',
                        session["cartons"],
                        f'{session["seconds_per_carton"]:.4f}',
                        session.get("name", ""),
                        session.get("goal", ""),
                        session.get("planned_minutes", ""),
                        f'{session.get("ema_seconds_per_carton", 0):.4f}',
                        session.get("full_target", ""),
                        session.get("session_id", ""),
                    ]
                )

        except Exception:
            pass

    # ========================================================
    # STYLE
    # ========================================================

    def apply_styles(self):

        self.setStyleSheet(
            f"""
            QMainWindow,
            QWidget#page,
            QStackedWidget {{
                background: {PAPER_2};
            }}

            QLabel {{
                font-family: {FONT};
            }}

            QFrame#sidebar {{
                background: {OLIVE_DARK};
                border-right: 3px solid {GOLD};
            }}

            QFrame#header {{
                background: {PAPER};
                border-bottom: 4px solid {GOLD};
            }}

            QFrame#card {{
                background: {PAPER_3};
                border: 1px solid {BORDER};
                border-radius: 14px;
            }}

            QFrame#timeBox {{
                background: {PAPER_3};
                border: 1px solid {BORDER};
                border-radius: 10px;
            }}

            QPushButton#navButton {{
                background: transparent;
                color: {PAPER};
                font-family: {FONT};
                font-size: 15px;
                font-weight: 700;
                border-radius: 10px;
                padding: 8px;
                text-align: center;
            }}

            QPushButton#navButton:hover {{
                background: {OLIVE};
            }}

            QPushButton#navButton:checked {{
                background: {GOLD};
                color: {NAVY_DARK};
            }}

            QFrame#sideFooter {{
                border-top: 1px solid {OLIVE};
            }}

            QLabel#sideTitle {{
                color: {PAPER};
                font-size: 20px;
                font-weight: 900;
            }}

            QLabel#sideSub {{
                color: #CFC6A8;
                font-size: 11px;
            }}

            QLabel#usbStatus {{
                color: #E9A58F;
                font-size: 11px;
                font-weight: 600;
            }}

            QLabel#usbStatus[connected="true"] {{
                color: #8FE0A0;
            }}

            QLabel#headerTitle {{
                color: {NAVY};
                font-size: 28px;
                font-weight: 900;
            }}

            QLabel#headerSub {{
                color: {OLIVE};
                font-size: 11px;
                font-weight: 800;
                letter-spacing: 2px;
            }}

            QLabel#verse {{
                color: {NAVY};
                font-size: 14px;
                font-weight: 600;
                font-style: italic;
            }}

            QLabel#motto {{
                color: {BROWN};
                font-size: 20px;
                font-weight: 900;
                border: 3px solid {BROWN};
                border-radius: 5px;
                padding: 2px 14px;
            }}

            QLabel#date {{
                color: {NAVY};
                font-size: 14px;
                font-weight: 700;
            }}

            QLabel#clock {{
                color: {NAVY};
                font-size: 26px;
                font-weight: 900;
            }}

            QLabel#homeTitle {{
                color: {NAVY};
                font-size: 22px;
                font-weight: 900;
            }}

            QLabel#bigCount {{
                color: {NAVY};
                background: {GOLD};
                border-radius: 16px;
                font-size: 120px;
                font-weight: 900;
                padding: 16px 22px;
            }}

            QLabel#goalLine {{
                color: {NAVY};
                font-size: 16px;
                font-weight: 800;
            }}

            QLabel#infoTitle {{
                color: {OLIVE};
                font-size: 14px;
                font-weight: 800;
            }}

            QLabel#infoValue {{
                color: {NAVY};
                font-size: 26px;
                font-weight: 900;
            }}

            QLabel#infoSubtitle {{
                color: {MUTED};
                font-size: 11px;
                font-weight: 700;
            }}

            QLabel#panelTitle {{
                color: {NAVY};
                font-size: 15px;
                font-weight: 900;
            }}

            QLabel#summary {{
                color: {BROWN};
                font-size: 13px;
                font-weight: 700;
            }}

            QPushButton[kind="green"] {{
                background: {GREEN};
                border-radius: 12px;
            }}

            QPushButton[kind="red"] {{
                background: {RED};
                border-radius: 12px;
            }}

            QPushButton[kind="blue"] {{
                background: {BLUE};
                border-radius: 12px;
            }}

            QPushButton[kind="orange"] {{
                background: {ORANGE};
                border-radius: 12px;
            }}

            QPushButton[kind]:hover {{
                border: 2px solid rgba(255,255,255,0.7);
            }}

            QPushButton[kind]:disabled {{
                background: #AAA49A;
            }}

            QLabel#actionTitle {{
                color: white;
                font-size: 16px;
                font-weight: 900;
            }}

            QLabel#actionSubtitle {{
                color: rgba(255,255,255,0.95);
                font-size: 11px;
                font-weight: 600;
            }}

            QLabel#pageTitle {{
                color: {NAVY};
                font-size: 30px;
                font-weight: 900;
            }}

            QLabel#pageSubtitle {{
                color: {OLIVE};
                font-size: 14px;
                font-weight: 700;
            }}

            QTableWidget {{
                background: {PAPER_3};
                alternate-background-color: #F0E4CB;
                border: 1px solid {BORDER};
                border-radius: 10px;
                gridline-color: #D7CAB5;
                font-family: {FONT};
                font-size: 12px;
                color: {NAVY_DARK};
                selection-background-color: #F1C96C;
                selection-color: {NAVY_DARK};
            }}

            QHeaderView::section {{
                background: {OLIVE};
                color: {PAPER_3};
                padding: 9px;
                border: none;
                font-family: {FONT};
                font-size: 12px;
                font-weight: 800;
            }}

            QDialog {{
                background: {PAPER_2};
            }}

            QPushButton#stepPlus,
            QPushButton#stepMinus {{
                color: white;
                border-radius: 10px;
                font-family: {FONT};
                font-size: 18px;
                font-weight: 900;
            }}

            QPushButton#stepPlus {{
                background: {BLUE};
            }}

            QPushButton#stepMinus {{
                background: {ORANGE};
            }}

            QPushButton#toolButton,
            QPushButton#dangerButton {{
                border-radius: 10px;
                padding: 6px 10px;
                font-family: {FONT};
                font-size: 14px;
                font-weight: 800;
            }}

            QPushButton#toolButton {{
                background: {GRAY};
                color: {NAVY};
                border: 1px solid {BORDER};
            }}

            QPushButton#dangerButton {{
                background: {PAPER_3};
                color: {RED};
                border: 2px solid {RED};
            }}

            QPushButton#stepPlus:hover,
            QPushButton#stepMinus:hover,
            QPushButton#toolButton:hover,
            QPushButton#dangerButton:hover {{
                border: 2px solid {NAVY};
            }}

            QPushButton#toolButton:disabled {{
                color: #9C9282;
            }}

            QLineEdit,
            QComboBox,
            QSpinBox {{
                background: white;
                border: 1px solid {BORDER};
                border-radius: 8px;
                padding: 8px 10px;
                color: {NAVY};
                font-family: {FONT};
                font-size: 14px;
                font-weight: 700;
            }}

            QCheckBox {{
                color: {NAVY_DARK};
                font-family: {FONT};
                font-size: 14px;
                font-weight: 700;
                spacing: 8px;
            }}

            QLabel#settingsTitle {{
                color: {NAVY};
                font-size: 18px;
                font-weight: 900;
            }}

            QLabel#settingsDesc {{
                color: {BROWN_SOFT};
                font-size: 12px;
                font-weight: 600;
            }}

            QPushButton#smallButton {{
                background: {GRAY};
                color: {NAVY};
                border: 1px solid {BORDER};
                border-radius: 8px;
                padding: 8px 12px;
                font-family: {FONT};
                font-size: 13px;
                font-weight: 800;
            }}

            QLabel#sensorState {{
                color: {RED};
                font-size: 16px;
                font-weight: 900;
            }}

            QPushButton#primaryButton {{
                background: {GREEN};
                color: white;
                border-radius: 10px;
                padding: 12px 22px;
                font-family: {FONT};
                font-size: 14px;
                font-weight: 800;
            }}

            QPushButton#secondaryButton {{
                background: {GRAY};
                color: {NAVY};
                border: 1px solid {BORDER};
                border-radius: 10px;
                padding: 12px 22px;
                font-family: {FONT};
                font-size: 14px;
                font-weight: 800;
            }}
            """
        )

    # ========================================================
    # CLOSE
    # ========================================================

    def closeEvent(
        self,
        event
    ):

        if self.session_active:

            res = QMessageBox.question(
                self,
                "Session Active",
                "Session is still active.\n"
                "End it and save before closing?",
                QMessageBox.Yes
                | QMessageBox.No
                | QMessageBox.Cancel
            )

            if res == QMessageBox.Cancel:

                event.ignore()

                return

            if res == QMessageBox.Yes:

                self.end_session()

        self.save_state()

        event.accept()


# ============================================================
# LOAD STATE
# ============================================================

def load_state():

    DATA_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    default = {
        "count": 0,
        "target": DEFAULT_TARGET,
        "work_period_minutes":
            DEFAULT_WORK_MINUTES,

        "serial":
            dict(DEFAULT_SERIAL),

        "last_session":
            None,

        "show_pie":
            True,
    }

    if not STATE_FILE.exists():

        return default

    try:

        data = json.loads(
            STATE_FILE.read_text(
                encoding="utf-8"
            )
        )

        default["count"] = max(
            0,
            int(
                data.get(
                    "count",
                    0
                )
            )
        )

        default["target"] = max(
            1,
            int(
                data.get(
                    "target",
                    DEFAULT_TARGET
                )
            )
        )

        default["work_period_minutes"] = max(
            1,
            int(
                data.get(
                    "work_period_minutes",
                    DEFAULT_WORK_MINUTES
                )
            )
        )

        if isinstance(
            data.get("serial"),
            dict
        ):

            default["serial"] = dict(
                DEFAULT_SERIAL,
                **data["serial"]
            )

        default["serial"]["port"] = FIXED_PORT
        default["serial"]["mode"] = FIXED_MODE

        default["last_session"] = (
            data.get(
                "last_session"
            )
        )

        default["show_pie"] = bool(
            data.get(
                "show_pie",
                True
            )
        )

        if isinstance(
            data.get("session_defaults"),
            dict
        ):

            default["session_defaults"] = (
                data["session_defaults"]
            )

    except Exception:
        pass

    return default


# ============================================================
# MAIN
# ============================================================

def main():

    app = QApplication(
        sys.argv
    )

    app.setApplicationName(
        f"{CAMP_NAME_AR} — "
        f"{CAMP_NAME} "
        f"{CAMP_TAGLINE}"
    )

    load_fonts()

    app.setFont(
        qfont(
            14,
            QFont.Normal
        )
    )

    state = load_state()

    dashboard = DashboardWindow(
        state
    )

    counter = CounterWindow(
        state
    )

    # ========================================================
    # SERIAL BRIDGE
    # ========================================================

    bridge = SerialBridge(
        state.get(
            "serial"
        )
    )

    # Debug output: run with python main.py from a terminal.
    bridge.raw_data_received.connect(
        lambda data: print(
            f"[SERIAL RAW] {datetime.now().strftime('%H:%M:%S.%f')[:-3]} | {data!r}",
            flush=True
        )
    )
    bridge.status_changed.connect(
        lambda text, connected: print(
            f"[SERIAL STATUS] connected={connected} | {text}", flush=True
        )
    )
    bridge.delta_received.connect(
        lambda delta: print(f"[SERIAL COUNT] delta={delta}", flush=True)
    )

    # Settings -> Serial Bridge
    dashboard.serial_config_changed.connect(
        bridge.set_config
    )

    # Sensor -> Dashboard
    bridge.delta_received.connect(
        dashboard.on_sensor_delta
    )

    # Sensor status -> Dashboard
    bridge.status_changed.connect(
        dashboard.on_sensor_status
    )

    # Sensor status -> Counter
    bridge.status_changed.connect(
        counter.set_status
    )

    # Dashboard -> Counter
    dashboard.count_changed.connect(
        lambda _c:
        counter.refresh()
    )

    dashboard.countdown_changed.connect(
        counter.set_countdown
    )

    dashboard.rate_changed.connect(
        counter.set_rate
    )

    counter.refresh()

    counter.set_countdown(
        0,
        False,
        False
    )

    # ========================================================
    # WINDOWS
    # ========================================================

    dashboard.show()

    screens = app.screens()

    if len(screens) >= 2:

        second_geo = (
            screens[1]
            .availableGeometry()
        )

        counter.setGeometry(
            second_geo
        )

        counter.showFullScreen()

    else:

        counter.showFullScreen()

    # ========================================================
    # SERIAL THREAD
    # ========================================================

    thread = threading.Thread(
        target=bridge.run,
        daemon=True
    )

    thread.start()

    # ========================================================
    # APP
    # ========================================================

    exit_code = app.exec()

    # ========================================================
    # CLEANUP
    # ========================================================

    bridge.stop()

    sys.exit(
        exit_code
    )


# ============================================================
# ENTRY
# ============================================================

if __name__ == "__main__":
    main()